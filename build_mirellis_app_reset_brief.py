from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path("deliverables/Mirellis App Reset — V1 Product Brief.docx")

INK = "2E2B28"
MUTED = "6E655D"
PAPER = "F6F1E8"
LINEN = "EFE7DB"
CLAY = "85776B"
RULE = "D8C6B1"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for idx, width in enumerate(widths):
        grid.gridCol_lst[idx].set(qn("w:w"), str(width))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=RULE):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(doc, text, style=None, size=11, color=INK, bold=False, italic=False, before=0, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    set_run(p.add_run(text), size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    set_run(p.add_run(text), size=10.5)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run(r, size={1: 16, 2: 13, 3: 11}[level], color=INK if level < 3 else MUTED, bold=(level < 3))
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    set_table_borders(table, color=RULE)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LINEN)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper() + "  ")
    set_run(r, size=9, color=CLAY, bold=True)
    r = p.add_run(text)
    set_run(r, size=10.5, color=INK)
    return table


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    set_table_widths(table, [2000, 7360])
    set_table_borders(table, color=RULE)
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], PAPER)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(label), size=9.5, color=MUTED, bold=True)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(value), size=10.5)
    return table


def add_flow_table(doc):
    table = doc.add_table(rows=2, cols=6)
    set_table_widths(table, [1560] * 6)
    set_table_borders(table, color=RULE)
    labels = ["Recognise", "Name", "Begin", "Notice", "Understand", "Choose"]
    details = [
        "Work has ended, but my mind has not.",
        "Choose the closest explicit description.",
        "One guided transition; no library.",
        "Person records whether anything shifted.",
        "Optional finding and context.",
        "Optional physical ritual tool.",
    ]
    for i, text in enumerate(labels):
        set_cell_shading(table.cell(0, i), PAPER)
        p = table.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), size=9.5, color=INK, bold=True)
        p = table.cell(1, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(details[i]), size=8.7, color=MUTED)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(11)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
header.paragraph_format.space_after = Pt(0)
set_run(header.add_run("MIRELLIS  /  PRODUCT RESET"), size=8.5, color=MUTED, bold=True)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_before = Pt(0)
set_run(footer.add_run("Internal working brief  •  August 2026"), size=8.5, color=MUTED)

# Cover / masthead
add_text(doc, "MIRELLIS", size=10, color=CLAY, bold=True, after=8)
title = add_text(doc, "App Reset — V1 Product Brief", size=25, color=INK, bold=True, after=4)
subtitle = add_text(doc, "A private, immediate experience for the transition a person is in right now.", size=13, color=MUTED, after=18)
add_key_value_table(doc, [
    ("Decision", "Approved direction: reset the app around one complete Work → Recovery experience."),
    ("Status", "Product brief approved. No live-site or app implementation is authorised by this document."),
    ("Canonical V1", "Sleep → Wakefulness; Wakefulness → Readiness; Readiness → Focus; Focus → Recovery."),
    ("First vertical slice", "Work → Recovery only."),
])

heading(doc, "The decision", 1)
add_callout(
    doc,
    "Product job",
    "Help a person move from the state they have explicitly named to one small, practical next action. The app is not a content library, a diagnosis tool, or a performance dashboard.",
)
add_text(doc, "Mirellis will preserve its institutional and research architecture, but change the consumer entry point. The person should not need to learn the model before receiving value.", after=8)

heading(doc, "Operating model", 1)
add_flow_table(doc)
add_text(doc, "The experience is intentionally ordered: practical value first, deeper explanation second, commerce last. The transition must work without a purchase.", size=10.5, color=MUTED, italic=True, before=8, after=4)

doc.add_page_break()
heading(doc, "What stays, moves, and stops", 1)
add_key_value_table(doc, [
    ("Stays", "The four canonical transitions, Observatory evidence and Findings, Sound as an experience ingredient, the optional physical products, local-first records, interruption/resume, and an optional personal reflection history."),
    ("Moves", "Mirror becomes a deliberate reflection path for people who want to explore patterns. It is not the default way into a known, immediate difficulty. Sound and ritual cease to be primary destinations and become parts of a transition."),
    ("Stops in V1", "Module-led home screens, generic Explore journeys, article libraries as the opening task, default assessments, scores, streaks, gamification, pressure to purchase, automatic research donation, and fake personalisation."),
    ("Adds", "A single recognition-led entry; one complete guided Work → Recovery transition; truthful outcome capture; optional links to the relevant Finding and Stillness Elixir after completion."),
])

heading(doc, "V1: Work → Recovery", 1)
add_text(doc, "Opening statement", size=10, color=CLAY, bold=True, after=2)
add_text(doc, "“Work has ended. Did your mind?”", size=14, color=INK, bold=True, after=10)
add_text(doc, "The first screen should offer recognition, not explanation. The person can begin directly from a website CTA or open the app and select this moment.", after=6)

heading(doc, "Entry and branch rules", 2)
add_text(doc, "Ask one question only: “What feels closest right now?”", size=10.5, bold=True, after=5)
add_bullet(doc, "“I’m still thinking about unfinished work.” → a concrete closure cue: name the next action and close the work surface.")
add_bullet(doc, "“I feel restless and keep checking things.” → an interruption cue: put the device away and shift attention through the selected sound environment.")
add_bullet(doc, "“I’m physically tired but mentally active.” → a settling cue: reduce light and follow a slower breath/tactile instruction.")
add_bullet(doc, "“I’m not sure; I just haven’t switched off.” → the simplest neutral version.")
add_callout(doc, "Branching constraint", "A choice may appear only when it produces a visibly different next instruction or experience. If the content does not differ, remove the choice rather than simulating personalisation.")

heading(doc, "The guided transition", 2)
add_number(doc, "Set the environment: one clear, non-medical instruction that can be completed without a Mirellis product.")
add_number(doc, "Begin the approved Work → Recovery sound environment. The interface becomes quiet; no article rail, competing navigation, or progress-performance language.")
add_number(doc, "Offer no more than a few paced prompts, matched to the person’s explicit choice. The practice should be completable in approximately 8–10 minutes.")
add_number(doc, "Support interruption and resume. Leaving early carries no penalty and never silently advances or saves a path.")

heading(doc, "Truthful close", 2)
add_text(doc, "Ask: “Did anything shift?”", size=10.5, bold=True, after=5)
add_bullet(doc, "Yes")
add_bullet(doc, "A little")
add_bullet(doc, "Not really")
add_text(doc, "Then offer an optional prompt: “What did you notice?” The note is the person’s own record; Mirellis does not interpret it.", after=6)
add_text(doc, "Primary action: “Back to today.” A saved practice or next-day suggestion may appear only if the person deliberately chooses an action whose visible consequence is stated.", size=10.5, color=MUTED, italic=True, after=8)

heading(doc, "After the experience", 1)
add_key_value_table(doc, [
    ("Research", "Optional: “Understand this transition.” Opens the related Observatory Finding or investigation only after the practice."),
    ("Product", "Optional: “A ritual tool for this transition.” Stillness Elixir is framed as a physical instrument, not a requirement or a cure."),
    ("Record", "A local observation record may be kept only when the person explicitly saves it. The person can edit, export, or delete it."),
    ("Research participation", "Never automatic. A separate, understandable consent request is required before any observation is used beyond the person’s own device."),
])

heading(doc, "Trust, privacy, and product boundaries", 1)
add_bullet(doc, "No scores, diagnoses, inferred readiness, hidden models, engagement optimisation, or social features.")
add_bullet(doc, "Every displayed observation is authored, selected, or confirmed by the person and linked to visible local evidence.")
add_bullet(doc, "No free-text interpretation. Mirellis may store a person’s note but must not infer meaning from it.")
add_bullet(doc, "The transition functions without purchase, account creation, or research participation.")
add_bullet(doc, "Accessible controls have a minimum 44 × 44 px target; keyboard focus is visible; save/completion feedback uses accessible status text; reduced motion is respected.")

heading(doc, "Minimal local-first record", 1)
add_key_value_table(doc, [
    ("Practice record", "transition_id, practice_id, selected_entry, started_at, completed_at or left_at, explicit_outcome, saved_by_person."),
    ("Reflection", "optional_note, authored_at, edited_at. No derived attributes, labels, or scores."),
    ("Consent", "separate research_consent value, timestamp, and clear scope. Default: no consent."),
    ("Person control", "view, edit, export, delete, and cancel deletion. Deletion confirmation must be cancelable."),
])

doc.add_page_break()
heading(doc, "Build order", 1)
add_bullet(doc, "First: audit the current app only to identify reusable infrastructure—sound delivery, local storage, outcome capture, resume handling, accessibility foundations, and deployment pipeline.")
add_bullet(doc, "Then: replace the consumer interface with the Work → Recovery vertical slice. Do not preserve module navigation for compatibility.")
add_bullet(doc, "Next: implement the four honest entry routes and their materially different first cues.")
add_bullet(doc, "Then: implement the practice, interruption/resume, truthful close, optional note, and local record.")
add_bullet(doc, "Afterward: add the post-completion Observatory and Stillness doors. They must be optional and secondary.")
add_bullet(doc, "Finally: test with real people before adding the second canonical transition.")

heading(doc, "Acceptance criteria for the first test", 1)
add_bullet(doc, "A first-time visitor can begin Work → Recovery in one decision from the landing screen.")
add_bullet(doc, "Every offered entry choice changes the next cue; there are no false branches.")
add_bullet(doc, "The practice works with no purchase and no account.")
add_bullet(doc, "People can leave, resume, and finish without losing agency or being silently moved to another practice.")
add_bullet(doc, "Outcome and note capture are optional, explicit, locally controlled, and accessible.")
add_bullet(doc, "The Observatory and product are available after the experience, not required before it.")
add_bullet(doc, "There are no scores, streaks, inferred progress, generic recommendations, fake Observatory content, or automatic research contribution.")

heading(doc, "Decisions still required before implementation", 1)
add_key_value_table(doc, [
    ("Practice content", "Approve the exact Work → Recovery script, sound asset, cue timing, and any product-use wording."),
    ("Outcome language", "Approve the final close copy and the exact choices; avoid language that implies recovery, cure, or measured improvement."),
    ("Research scope", "Decide whether V1 includes a separate, optional research-consent flow or remains fully local-only for the first test."),
    ("Commercial door", "Confirm whether Stillness Elixir appears in V1 after completion or is deferred until usefulness is demonstrated."),
    ("Evaluation", "Approve a small, consent-respecting test plan focused on starts, completion, self-reported outcome, research opens, and voluntary product exploration—not engagement optimisation."),
])

add_callout(
    doc,
    "Version freeze",
    "V1 is one complete Work → Recovery experience. Do not add new transitions, a dashboard, a content library, scores, AI interpretation, or broader personalisation until the first vertical slice has been tested with real people and deliberately reviewed.",
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
