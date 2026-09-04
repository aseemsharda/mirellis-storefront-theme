from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path("deliverables/Mirellis App — Unified V1 + V2 Pre-Build Specification v1.0.docx")

# compact_reference_guide preset with one named Mirellis warm-governance color override.
FONT = "Calibri"
INK = "2E2B28"
MUTED = "6E655D"
ACCENT = "766457"
ACCENT_DARK = "4E4037"
PAPER = "F7F2EA"
LINEN = "EEE5D8"
PALE = "F4F1EC"
RULE = "D8C6B1"
WHITE = "FFFFFF"
RED = "8A3232"
GREEN = "315B4A"
GOLD = "7A5A00"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
LAST_WAS_NUMBER = False
ACTIVE_NUM_ID = None


def reset_number_sequence():
    global LAST_WAS_NUMBER, ACTIVE_NUM_ID
    LAST_WAS_NUMBER = False
    ACTIVE_NUM_ID = None


def set_run(run, size=11, color=INK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    assert sum(widths) == PAGE_WIDTH_DXA, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    while len(grid.gridCol_lst) < len(widths):
        grid.append(OxmlElement("w:gridCol"))
    for index, width in enumerate(widths):
        grid.gridCol_lst[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=RULE, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def style_cell_text(cell, size=9.2, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        for run in p.runs:
            set_run(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths, font_size=9.0, header_fill=LINEN):
    reset_number_sequence()
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        p.add_run(str(header))
        style_cell_text(cell, size=8.8, color=ACCENT_DARK, bold=True)
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for c_index, value in enumerate(row):
            if r_index % 2 == 1:
                shade_cell(cells[c_index], PALE)
            p = cells[c_index].paragraphs[0]
            p.add_run(str(value))
            style_cell_text(cells[c_index], size=font_size)
    spacer(doc, 7)
    return table


def paragraph(doc, text="", size=11, color=INK, bold=False, italic=False, before=0, after=6,
              align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    reset_number_sequence()
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep
    if text:
        set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def rich_paragraph(doc, parts, before=0, after=6, size=11, color=INK):
    p = paragraph(doc, before=before, after=after)
    for part in parts:
        if isinstance(part, str):
            text, bold, italic, part_color = part, False, False, color
        else:
            text = part.get("text", "")
            bold = part.get("bold", False)
            italic = part.get("italic", False)
            part_color = part.get("color", color)
        set_run(p.add_run(text), size=size, color=part_color, bold=bold, italic=italic)
    return p


def heading(doc, text, level=1):
    reset_number_sequence()
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def bullet(doc, text, level=0):
    reset_number_sequence()
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), size=10.5)
    return p


def new_numbering_instance(doc, style_name):
    numbering = doc.part.numbering_part.element
    style = doc.styles[style_name]._element
    style_num = style.find(".//" + qn("w:numId"))
    base_num_id = style_num.get(qn("w:val")) if style_num is not None else "5"
    base_num = None
    for node in numbering.findall(qn("w:num")):
        if node.get(qn("w:numId")) == base_num_id:
            base_num = node
            break
    abstract_id = "0"
    if base_num is not None:
        abstract = base_num.find(qn("w:abstractNumId"))
        if abstract is not None:
            abstract_id = abstract.get(qn("w:val"))
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_id = max(existing or [0]) + 1
    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    new_num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    new_num.append(override)
    numbering.append(new_num)
    return new_id


def assign_numbering(paragraph_obj, num_id, level=0):
    p_pr = paragraph_obj._p.get_or_add_pPr()
    old = p_pr.find(qn("w:numPr"))
    if old is not None:
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)


def number(doc, text, level=0):
    global LAST_WAS_NUMBER, ACTIVE_NUM_ID
    style = "List Number" if level == 0 else "List Number 2"
    if not LAST_WAS_NUMBER or ACTIVE_NUM_ID is None:
        ACTIVE_NUM_ID = new_numbering_instance(doc, style)
    p = doc.add_paragraph(style=style)
    assign_numbering(p, ACTIVE_NUM_ID, level=level)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text), size=10.5)
    LAST_WAS_NUMBER = True
    return p


def spacer(doc, points=6):
    reset_number_sequence()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(points)
    return p


def callout(doc, label, text, fill=PAPER, label_color=ACCENT, body_color=INK):
    reset_number_sequence()
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), RULE)
        borders.append(border)
    p_pr.append(borders)
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.18
    set_run(p.add_run(label.upper() + "  "), size=8.8, color=label_color, bold=True)
    set_run(p.add_run(text), size=10.4, color=body_color)
    return p


def key_value(doc, rows):
    return add_table(doc, ["Decision", "Definition"], rows, [2000, 7360], font_size=9.5)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_DARK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def source_line(doc, number_value, title, url, note):
    p = paragraph(doc, before=0, after=4, size=9.5)
    set_run(p.add_run(f"{number_value}. "), size=9.5, color=ACCENT_DARK, bold=True)
    add_hyperlink(p, title, url)
    set_run(p.add_run(f" — {note}"), size=9.5, color=MUTED)


def page_break(doc):
    reset_number_sequence()
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

# Explicit compact_reference_guide style tokens.
normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for level, size, before, after, color in (
    (1, 16, 18, 10, ACCENT_DARK),
    (2, 13, 14, 7, ACCENT),
    (3, 12, 10, 5, MUTED),
):
    style = doc.styles[f"Heading {level}"]
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
    style = doc.styles[style_name]
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

# Running furniture.
header = section.header.paragraphs[0]
header.paragraph_format.space_after = Pt(0)
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run(header.add_run("MIRELLIS  /  UNIFIED APP SPECIFICATION"), size=8.2, color=MUTED, bold=True)
footer = section.footer.paragraphs[0]
footer.paragraph_format.space_before = Pt(0)
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(footer.add_run("Internal pre-build document  •  Page "), size=8.2, color=MUTED)
add_page_field(footer)

# Memo-masthead opening.
spacer(doc, 10)
paragraph(doc, "MIRELLIS · RITUAL SCIENCE", size=9.5, color=ACCENT, bold=True, after=8)
paragraph(doc, "Unified V1 + V2 Pre-Build Specification", size=25, color=INK, bold=True, after=4)
paragraph(doc, "Product, clinical content, experience, trust, data and engineering authority", size=13, color=MUTED, after=18)
key_value(doc, [
    ("Version", "1.0 — review draft for approval before implementation"),
    ("Date", "19 August 2026"),
    ("Applies to", "app.mirellis.in: current V1 After Work experience and proposed V2 capabilities"),
    ("Implementation status", "No code or live deployment is authorised by this document."),
    ("Primary product decision", "Mirellis becomes an action-first, private companion for four difficult moments across the day."),
])
callout(doc, "Governing outcome", "A person should leave Mirellis having done something real in their life—not merely having read advice, completed a quiz, or received a label.")

heading(doc, "How to use this document", 1)
paragraph(doc, "This is the implementation gate. It converts the research pack and earlier product authorities into one buildable system. Consumer copy is written exactly where marked. Internal identifiers exist only for engineering and content governance and must never appear on consumer screens.")
bullet(doc, "Approve or amend the product decisions and open decisions in Section 22.")
bullet(doc, "Clinically review the safety routes and Tier C exclusions before implementation.")
bullet(doc, "Build only in the phased order in Section 19, behind a preview environment and feature flags.")
bullet(doc, "Do not deploy to production until the acceptance tests in Section 21 pass.")

heading(doc, "Contents", 1)
for item in [
    "1. Governing authorities and decision hierarchy",
    "2. Executive product decisions",
    "3. Canonical V1 model and V2 relationship",
    "4. Information architecture and consumer journeys",
    "5. Shared action engine and personalisation",
    "6. Common interface and interaction specification",
    "7. Waking experience",
    "8. Focus experience",
    "9. After Work and Restoration experience",
    "10. Night Audit and sleep-support experience",
    "11. Launch intervention library",
    "12. Motivation, reminders and retention without manipulation",
    "13. Personal Observatory integration",
    "14. Practitioner directory and safety routing",
    "15. Data model, privacy, consent and deletion",
    "16. AI boundaries and content governance",
    "17. Engineering architecture, PWA/native boundaries and cost",
    "18. Measurement and learning operations",
    "19. Build order and release plan",
    "20. Accessibility and inclusive design",
    "21. Test plan and acceptance criteria",
    "22. Approval decisions and version freeze",
    "Appendix A. Event and object dictionaries",
    "Appendix B. Evidence register",
]:
    bullet(doc, item)


page_break(doc)
heading(doc, "1. Governing authorities and decision hierarchy", 1)
paragraph(doc, "This specification consolidates product decisions; it does not erase the earlier trust and research authorities. Where older consumer-flow instructions conflict with this specification, this specification governs the V1 + V2 application experience. Earlier privacy, evidence, consent and stewardship constraints continue to govern.")
add_table(doc, ["Authority", "What remains binding", "Relationship to this specification"], [
    ("Lambda Constitution", "Human agency, truthfulness, restraint and non-manipulation.", "Highest authority; cannot be overridden by growth or engagement goals."),
    ("V1 Transition Platform", "Four canonical V1 transitions and interruption/resume foundations.", "Preserved internally; consumer entry becomes problem-first."),
    ("V2 Learning Intelligence Layer", "Explicit evidence, user confirmation and bounded learning.", "Implemented without hidden inference or unconfirmed personal labels."),
    ("V3 Curriculum Stewardship", "Versioning, evidence review, content ownership and change control.", "Applies to every intervention object and public claim."),
    ("Launch & Learning Operations", "Controlled rollout, observed failures and version freeze.", "Operational basis for feature flags and staged release."),
    ("Personal Observatory v4", "Private mirror; user-authored, selected or confirmed observations only.", "Governs personal records, timeline, export and deletion."),
    ("App Reset V1 Product Brief", "Immediate value, no fake branching, Work → Recovery vertical slice.", "Retained and expanded into a unified four-moment system."),
], [1700, 3530, 4130], font_size=8.7)

heading(doc, "Decision hierarchy", 2)
number(doc, "Safety and truthful claims override conversion, completion and retention.")
number(doc, "A user’s explicit choice overrides system convenience.")
number(doc, "A real-world action overrides additional content consumption.")
number(doc, "Deterministic, reviewed routing overrides generated recommendations.")
number(doc, "If evidence is weak, the app must say so through restrained copy and optional positioning.")
number(doc, "If a branch does not change the next action, remove the branch.")

heading(doc, "Non-negotiable prohibitions", 2)
for item in [
    "No diagnosis, clinical verdict, mental-health interpretation or inferred personality model.",
    "No nervous-system, sleep, focus, readiness, recovery or wellbeing score.",
    "No streaks, leaderboards, public comparison, competitive ranking or loss-framed reminders.",
    "No automatic research donation and no reuse of private notes for model training.",
    "No free-text interpretation in the launch release.",
    "No supplement, medication, sleep-restriction or chronotherapy protocol without separate clinical governance.",
    "No claim that a single exercise permanently rewires attention, guarantees sleep or fixes a circadian disorder.",
    "No internal identifiers, transition arrows, module labels, levels or curriculum terminology on consumer screens.",
]:
    bullet(doc, item)


heading(doc, "2. Executive product decisions", 1)
callout(doc, "Product definition", "Mirellis is a private action companion for moments when a person knows what they want to do—wake, focus, stop work or stop for sleep—but cannot reliably cross the boundary alone.")
add_table(doc, ["Question", "Decision"], [
    ("What does the app sell first?", "No subscription at launch. It proves usefulness through free, complete experiences."),
    ("What is the primary value?", "A fitting action performed now, followed by a truthful personal record."),
    ("Is every journey a course?", "No. Most visits are immediate sessions; audits and programs appear only when they fit."),
    ("Is V2 a replacement?", "No. It adds a shared engine and new capabilities while preserving V1 progress and records."),
    ("Is AI required?", "No. V2 launch must work without AI. AI may later explain approved options within strict boundaries."),
    ("Does the app diagnose?", "No. It identifies user-confirmed patterns and routes concerning or persistent experiences to qualified care."),
    ("How is motivation created?", "Clear appointments, real-world precommitment, reduced friction, immediate feedback and remembered context—not pressure."),
    ("How is personalisation created?", "A governed intervention library plus deterministic eligibility rules, not thousands of generated answers."),
], [2500, 6860], font_size=9.2)

heading(doc, "The product promise", 2)
paragraph(doc, "Consumer-facing promise: “Tell Mirellis what is difficult right now. It will help you choose and begin one practical next action.”")
paragraph(doc, "The promise is deliberately narrower than “improve your sleep, focus and restoration.” Mirellis may help a person test what works for them, but it does not promise treatment, transformation or optimisation.")

heading(doc, "The five experience modes", 2)
add_table(doc, ["Mode", "Time", "Use", "End state"], [
    ("Help me now", "2–10 min", "One immediate difficulty and one practical action.", "Action started or consciously declined."),
    ("Understand the pattern", "3–7 min", "Short reconstruction when the right action is unclear.", "User confirms, rejects or postpones a candidate explanation."),
    ("Run an experiment", "3 sessions to 10 days", "Test one change with a defined comparison.", "User sees their own before/after evidence without a score."),
    ("Build a practice", "Up to 21 sessions", "Optional maintenance after a useful action is identified.", "Reusable routine the person can perform without Mirellis."),
    ("Find professional help", "Immediate", "Persistent, severe, safety-relevant or out-of-scope concerns.", "Clear next step and directory filters; no diagnosis."),
], [1540, 1120, 3720, 2980], font_size=8.7)


heading(doc, "3. Canonical V1 model and V2 relationship", 1)
paragraph(doc, "The canonical V1 curriculum remains four transitions. V2 does not imply seven transitions and does not invent a new curriculum. It changes how people enter the system and adds a Night Audit as a cross-cutting capability.")
add_table(doc, ["Canonical V1 transition", "Consumer moment", "V2 capability"], [
    ("Sleep → Wakefulness", "Getting going after waking", "Waking support and recurring-pattern experiment."),
    ("Wakefulness → Readiness", "Preparing to begin", "Start-friction and task-clarity support."),
    ("Readiness → Focus", "Beginning and holding focus", "Real work session, interruption capture and adaptation."),
    ("Focus → Recovery", "Leaving work behind", "V1 After Work practice plus restoration routes."),
], [2550, 3200, 3610], font_size=9.0)
callout(doc, "Night Audit status", "Night Audit is a V2 diagnostic-and-experiment tool connected to sleep and restoration. It is not displayed as a fifth V1 transition and does not change the locked four-transition curriculum.")

heading(doc, "The 24-hour continuity model", 2)
paragraph(doc, "Mirellis treats the day as a linked system without claiming that one domain causes another in every person. The app may invite the person to test a connection, but it must not infer one.")
add_table(doc, ["If the person explicitly reports…", "Mirellis may offer…", "Mirellis must not say…"], [
    ("Work remains mentally open at night", "The After Work closure action before a Night experiment.", "“Your sleep problem is caused by work.”"),
    ("Insufficient sleep before a focus session", "A capacity-protection session or rescheduling option.", "“You cannot focus because you are sleep deprived.”"),
    ("They are not sleepy at intended bedtime", "Timing observation and professional information if persistent.", "“You have a delayed circadian phase.”"),
    ("Morning grogginess persists unusually long", "Safety advice and an option to seek assessment.", "“You have a sleep disorder.”"),
], [3020, 3410, 2930], font_size=8.8)


heading(doc, "4. Information architecture and consumer journeys", 1)
heading(doc, "Primary navigation", 2)
add_table(doc, ["Destination", "Job", "Launch behavior"], [
    ("Today", "One contextual entry and the active or recommended action.", "Default tab. Never opens with a catalogue."),
    ("Practices", "Deliberate exploration of available immediate sessions, experiments and saved routines.", "Grouped by familiar difficulties, not transition IDs."),
    ("My record", "Private Observatory: only deliberate saves and completed experiences.", "Shows a true empty state before meaningful use."),
    ("Help", "Safety information, practitioner directory and product boundaries.", "Available from every route without abandoning saved state."),
    ("Settings", "Privacy, reminders, accessibility, export and deletion.", "No analytics toggle if analytics are not active."),
], [1450, 4210, 3700], font_size=8.8)

heading(doc, "Today entry", 2)
paragraph(doc, "Exact headline: “What is difficult right now?”", bold=True)
add_table(doc, ["Visible choice", "Internal mapping", "Immediate next question"], [
    ("Getting going after waking", "Waking", "“What is hardest right now?”"),
    ("Beginning or holding focus", "Readiness / Focus", "“What must exist when this session ends?”"),
    ("Leaving work behind", "After Work / Restoration", "“What is still with you?”"),
    ("Stopping for sleep", "Night Audit", "“What happened when you meant to stop?”"),
    ("I’m not sure", "Neutral orientation", "“Which moment are you closest to: waking, working, finishing work or stopping for sleep?”"),
], [2890, 2650, 3820], font_size=8.8)

heading(doc, "First-time journey", 2)
number(doc, "Land on Today. Read one sentence explaining that Mirellis helps with difficult boundaries in the day.")
number(doc, "Choose the current difficulty; no account, subscription or transition education is required.")
number(doc, "Answer only the questions necessary to select a safe action.")
number(doc, "See one recommendation with: why it fits, what will happen, duration and the small promise.")
number(doc, "Begin the action. The app becomes visually quiet and avoids unrelated navigation.")
number(doc, "Record only a truthful outcome. Saving a note or practice is optional and explicit.")
number(doc, "Return to Today, run a relevant experiment, choose another approach, or leave for now.")

heading(doc, "Returning journey", 2)
bullet(doc, "If an action is unfinished: “Continue [practice name].”")
bullet(doc, "If an action was completed: “Return to [practice name].”")
bullet(doc, "If an action was deliberately saved: show its visible consequence and offer it directly.")
bullet(doc, "If nothing is active: ask the current problem again; do not assume yesterday’s state persists.")
bullet(doc, "Never call an untouched practice “Continue.”")


heading(doc, "5. Shared action engine and personalisation", 1)
paragraph(doc, "The engine has six deterministic stages. Some immediate sessions skip reconstruction and move directly from recognition to action.")
add_table(doc, ["Stage", "System behavior", "User control"], [
    ("Recognise", "Capture the current difficulty and available time.", "Choose, skip or leave."),
    ("Reconstruct", "Ask the smallest set of behavioral questions needed.", "Answer approximately, select “not sure,” or stop."),
    ("Confirm", "Present one or two candidate patterns as possibilities.", "Confirm, reject or say neither fits."),
    ("Act", "Offer one reviewed action and honest alternatives.", "Begin, choose another, or leave."),
    ("Record", "Store timestamps and explicit outcomes only.", "Save, edit, discard or keep local only."),
    ("Adapt", "Use explicit history to change the next eligible action.", "See why it changed and override it."),
], [1450, 4880, 3030], font_size=8.8)

heading(doc, "Recommendation contract", 2)
paragraph(doc, "Every recommendation card must contain all five elements:")
number(doc, "Acknowledgement: repeat the difficulty in human language.")
number(doc, "Fit: explain which explicit answer made this action relevant.")
number(doc, "Small promise: state what the action is intended to make easier—never what it will guarantee.")
number(doc, "Effort: exact duration, required environment and whether the phone must be put down.")
number(doc, "Choice: “Start this practice,” “Choose another,” or “Leave for now.”")

heading(doc, "Candidate-pattern rules", 2)
bullet(doc, "A candidate pattern is displayed only when all required explicit signals are present.")
bullet(doc, "Use possibility language: “One possibility is…” or “This may fit because…”")
bullet(doc, "Never display a pattern inferred from free text.")
bullet(doc, "The person must confirm the pattern before it becomes part of My record.")
bullet(doc, "If two patterns remain plausible, show both without ranking them as stronger or weaker.")
bullet(doc, "If no governed route fits, say so and provide neutral alternatives or professional support.")

heading(doc, "Intervention-object contract", 2)
add_table(doc, ["Field", "Required definition"], [
    ("internal_key", "Stable engineering identifier; never consumer-visible."),
    ("consumer_name", "Plain-language action name."),
    ("problem_tags", "Explicit difficulties the action may address."),
    ("eligibility", "Answers and context required before recommendation."),
    ("exclusions", "Safety, capability or evidence reasons not to offer."),
    ("instruction", "Exact action; no vague advice."),
    ("duration", "Expected elapsed and active time."),
    ("evidence_tier", "A, B or C under Section 11."),
    ("claim_text", "Approved, qualified explanation of why it may help."),
    ("completion_evidence", "What the user explicitly records afterward."),
    ("next_options", "Honest actions available at completion."),
    ("version", "Content version and clinical-review status."),
], [2200, 7160], font_size=9.1)


heading(doc, "6. Common interface and interaction specification", 1)
heading(doc, "Visual and behavioral direction", 2)
bullet(doc, "Mobile-first application shell; it must not resemble a document, article page or Google Form.")
bullet(doc, "One primary card per screen, one dominant action, restrained secondary choices and persistent access to leave.")
bullet(doc, "Warm-neutral Mirellis palette, strong typography, generous spacing and a clear active state.")
bullet(doc, "State change is communicated through restrained light, spacing, contrast and motion; reduced-motion mode removes nonessential animation.")
bullet(doc, "During an off-screen action, the interface becomes a minimal timer or confirmation surface rather than serving more content.")

heading(doc, "Screen anatomy", 2)
add_table(doc, ["Region", "Requirement"], [
    ("Header", "Mirellis brand, current destination and one settings/help affordance. No collision on 320 px-wide screens."),
    ("Context line", "Where the person is and why this screen exists; no internal sequence labels."),
    ("Primary card", "One prompt, recommendation or active action."),
    ("Action area", "Primary control first; alternatives visibly subordinate but usable."),
    ("Status area", "aria-live feedback for saves, completion, reminder changes and errors."),
    ("Exit", "“Leave for now” remains reachable and never carries a penalty."),
], [1850, 7510], font_size=9.0)

heading(doc, "Progress language", 2)
add_table(doc, ["State", "Required copy", "Prohibited copy"], [
    ("Untouched", "Start with… / Start this practice", "Continue / Resume"),
    ("Unfinished", "Continue…", "Start again / Failed"),
    ("Completed", "Return to…", "Continue / Next level"),
    ("Saved", "Saved. [Visible consequence].", "Saved, with no explanation"),
], [1800, 3400, 4160], font_size=9.1)

heading(doc, "Universal completion", 2)
paragraph(doc, "For a first completed action, use:")
callout(doc, "Exact close", "Heading: “You made a first choice.” Body: “That is enough for today.” Primary action: “Back to today.”")
paragraph(doc, "“Keep this practice for tomorrow” may appear only if it pins the practice to Tomorrow and immediately shows: “Saved. Tomorrow’s suggestion will build on this practice.”")

heading(doc, "Permissions and reminders", 2)
bullet(doc, "Explain the mechanism before requesting permission: browser notification, calendar file, email or future native notification.")
bullet(doc, "Never describe a calendar file as an alarm or guarantee that a PWA notification will fire when the browser or operating system prevents it.")
bullet(doc, "Ask for notification permission only after the person explicitly schedules something.")
bullet(doc, "Every reminder has edit, snooze and disable controls. Missing a reminder changes no score or status.")


page_break(doc)
heading(doc, "7. Waking experience", 1)
callout(doc, "Consumer job", "Help me take an appropriate first step after waking without pretending that all morning difficulty is the same or that a brief routine can replace adequate sleep or assessment.")

heading(doc, "Entry copy", 2)
paragraph(doc, "Headline: “You’re awake. What is making it hard to begin?”", bold=True)
paragraph(doc, "Support: “Mirellis will help you choose one first action. It will not judge how well you slept.”")

heading(doc, "Exact question set", 2)
add_table(doc, ["Key", "Exact prompt", "Response control", "When asked"], [
    ("WK-Q1", "What is hardest right now?", "Single choice: Sleepy or foggy / Physically heavy / Anxious or rushed / I do not know where to begin / I feel unwell or unsafe / Not sure", "Always"),
    ("WK-Q2", "Roughly how much opportunity did you have to sleep?", "Single choice: Under 5 hours / 5–6 / 6–7 / 7–9 / More than 9 / Not sure or prefer not to say", "If sleepy, foggy or heavy"),
    ("WK-Q3", "How long have you been awake?", "Single choice: Under 10 minutes / 10–30 / 30–90 / More than 90 / Not sure", "If sleepy, foggy or heavy"),
    ("WK-Q4", "Do you need to drive, operate machinery or make a safety-critical decision within the next hour?", "Yes / No", "If sleepy, foggy, heavy or unsafe"),
    ("WK-Q5", "Has this level of morning difficulty been happening most days for two weeks or longer?", "Yes / No / Not sure", "If more than 90 minutes awake or user chooses recurring help"),
    ("WK-Q6", "What can you access right now?", "Multiple choice: Outdoor daylight / Indoor light / Space to move / A quiet minute / None of these", "Before action selection"),
], [900, 3780, 3120, 1560], font_size=8.4)

heading(doc, "Routing rules", 2)
add_table(doc, ["Confirmed situation", "Required first recommendation", "Alternatives / boundary"], [
    ("Recently awake; sleepy/foggy; adequate or unknown sleep; no safety flag", "Open light and movement: move to available bright/outdoor light and complete gentle movement for 2–5 minutes.", "Quiet orientation if light/movement unavailable. State that this may support alertness; do not promise performance."),
    ("Short sleep opportunity", "Protect limited capacity: choose one low-risk first task and postpone an avoidable high-stakes decision if possible.", "Offer a short orientation action. Do not prescribe compensatory caffeine or claim recovery."),
    ("Anxious or rushed", "Orient before accelerating: identify where you are, what must happen first and one action that can wait.", "Optional 60–90 second paced exhale or stillness audio."),
    ("No clear start", "Name the first visible action for the morning and begin it for two minutes.", "Choose another first action; no productivity score."),
    ("Safety-critical task + significant sleepiness", "Safety route: advise against relying on a quick exercise to restore safe performance and encourage an appropriate safer plan.", "Offer help information; do not start a performance challenge."),
    ("Persistent or concerning morning sleepiness", "Explain that persistent sleepiness can have multiple causes and offer professional assessment information.", "May still offer a gentle immediate action if the person chooses; never label a disorder."),
], [2900, 3900, 2560], font_size=8.4)

heading(doc, "Immediate action scripts", 2)
heading(doc, "Open light and movement", 3)
number(doc, "Move toward the brightest comfortable light available; outdoors is offered only if practical and safe.")
number(doc, "Stand or move gently for two minutes. The person chooses walking, mobility or simple household movement.")
number(doc, "Name one first task. Mirellis offers to begin it or end the session.")
paragraph(doc, "Approved fit copy: “You said you have only recently woken and feel foggy. Light and movement may support alertness, but this is not a substitute for enough sleep.”", size=10, color=MUTED, italic=True)

heading(doc, "Orient before accelerating", 3)
number(doc, "Show the current time and the person’s first fixed commitment if they entered one.")
number(doc, "Ask: “What truly has to happen first?” Provide a one-line response field and three optional templates.")
number(doc, "Ask: “What can wait?” Save only if the person explicitly chooses to keep the plan.")

heading(doc, "Waking experiment", 2)
paragraph(doc, "Offer only after three separate Waking visits or a deliberate request. Duration: seven mornings. The person chooses one stable anchor to test: wake-time consistency, immediate light opportunity, or a two-minute first-action routine. Mirellis records the chosen action and explicit morning outcome; it does not calculate a readiness score.")
add_table(doc, ["Morning record", "Allowed value"], [
    ("Action attempted", "Yes / Partly / No"),
    ("Beginning felt", "Easier / About the same / Harder / Not sure"),
    ("Context", "Approximate sleep opportunity and time since waking, if provided"),
    ("Personal note", "Optional user-authored text; never interpreted"),
], [2400, 6960], font_size=9.2)


heading(doc, "8. Focus experience", 1)
callout(doc, "Consumer job", "Help me begin and complete a real piece of work by changing the task and environment—not by asking me to play a generic attention game.")

heading(doc, "Entry copy", 2)
paragraph(doc, "Headline: “What must exist when this session ends?”", bold=True)
paragraph(doc, "Support: “Choose a real result. Mirellis will help you prepare the conditions, then get out of the way.”")

heading(doc, "Exact question set", 2)
add_table(doc, ["Key", "Exact prompt", "Response control", "When asked"], [
    ("FC-Q1", "What must exist when this session ends?", "One-line user response or template: A draft / A decision / A completed section / A reviewed set / A clear next step", "Always"),
    ("FC-Q2", "What is stopping you right now?", "Single choice: I cannot begin / The task is unclear / My phone or sites pull me away / Other thoughts keep arriving / People or messages interrupt me / I am too tired / Not sure", "Always"),
    ("FC-Q3", "How much uninterrupted time do you actually have?", "10 / 20 / 35 / 50 minutes / Choose another time", "Always"),
    ("FC-Q4", "Where is your phone?", "Another room / In a bag or drawer / On the desk / I need it for this task / Prefer not to say", "If digital pull or first session"),
    ("FC-Q5", "Which tools do you genuinely need?", "Multiple selection of user-entered apps/sites or No digital tools", "If digital pull"),
    ("FC-Q6", "What is the first visible action?", "One-line response; system offers a concrete rewrite for confirmation only", "If cannot begin or task unclear"),
    ("FC-Q7", "Is your tiredness strong enough that continuing would be unsafe or clearly unproductive?", "Yes / No / Not sure", "If too tired"),
], [900, 3690, 3210, 1560], font_size=8.3)

heading(doc, "Focus route rules", 2)
add_table(doc, ["Explicit barrier", "Preparation", "Active session"], [
    ("Cannot begin", "Convert the outcome into one visible physical or digital action; use a brief countdown.", "Start with a 5-minute launch inside the chosen session length."),
    ("Task unclear", "Define the deliverable, done condition and first action.", "Work only toward that deliverable; park other tasks."),
    ("Phone/sites", "Relocate phone; close nonrequired tabs; define needed tools. Native blocking offered only when supported.", "Minimal timer plus one-tap interruption capture after the session—not during every distraction."),
    ("Other thoughts", "Create a temporary parking line for thoughts without interpreting them.", "Return to the chosen result; review the parking line after the session."),
    ("People/messages", "Choose a visible boundary: status message, closed door, headphones or scheduled return.", "Session length must fit the real protected window."),
    ("Too tired", "Choose a shorter low-risk task, a movement break, reschedule or stop.", "Never frame pushing through as discipline; safety route if required."),
    ("Not sure", "Use the simplest preparation: result, first action, phone position and available time.", "Run one session; ask what actually interrupted afterward."),
], [2230, 4160, 2970], font_size=8.4)

heading(doc, "Focus session sequence", 2)
number(doc, "Result card: repeat the user’s intended output and chosen duration.")
number(doc, "Preparation card: show no more than three actions, each requiring confirmation. Example: “Phone in another room,” “Only the document and research tab open,” “First action: write the three section headings.”")
number(doc, "Commitment card: “For the next 20 minutes, I am working only toward [result].” The user may edit or decline.")
number(doc, "Active state: timer, pause for a real interruption, end early and an unobtrusive link to the parking line. No content feed.")
number(doc, "Close: record the observable result and interruption source. Do not grade the session.")

heading(doc, "Exact focus close", 2)
paragraph(doc, "Question 1: “Did you begin the task you chose?” — Yes / Partly / No", bold=True)
paragraph(doc, "Question 2: “What exists now?” — The result I planned / A useful part of it / A clearer next step / Something different but useful / Nothing visible yet", bold=True)
paragraph(doc, "Question 3: “What interrupted you most?” — Nothing / Phone / Websites / Messages or people / Another task / My thoughts / Fatigue / Something else", bold=True)
paragraph(doc, "Completion copy: “You created [user-selected result]. Mirellis will use only what you chose here if you ask for another session.”")

heading(doc, "Focus experiment", 2)
paragraph(doc, "Duration: three real work sessions, not three calendar days. Session 1 establishes the observed barrier. Sessions 2 and 3 test one controlled change, such as phone location, first-action definition or a protected communication boundary. The comparison is behavior-based: began or did not begin; result produced; interruption selected. No focus score or matched-decision percentage is calculated.")

heading(doc, "PWA and native enforcement", 2)
bullet(doc, "PWA launch: preparation checklist, browser full-screen session, optional browser notification, calendar reminder and explicit phone-location commitment.")
bullet(doc, "Native later: operating-system app blocking, focus filters, reliable notifications, widgets and background state where platform policy permits.")
bullet(doc, "Never claim an app is blocked unless the operating system confirms that blocking is active.")


heading(doc, "9. After Work and Restoration experience", 1)
callout(doc, "Consumer job", "Help me make work feel finished enough to enter the evening, using a route that matches what is actually still with me.")

heading(doc, "Entry copy", 2)
paragraph(doc, "Headline: “Work has ended. What is still with you?”", bold=True)
paragraph(doc, "Support: “Choose the closest description. Each answer leads to a different first action.”")

heading(doc, "Exact question set", 2)
add_table(doc, ["Key", "Exact prompt", "Response control", "When asked"], [
    ("AW-Q1", "What is still with you?", "Single choice: An unfinished task / I keep checking / My mind or body is still activated / Something at work affected me emotionally / I feel depleted or numb / I cannot tell", "Always"),
    ("AW-Q2", "Where are you now?", "At the workspace / Commuting / At home / Somewhere else", "Always"),
    ("AW-Q3", "How much time can you give this transition?", "3 / 8 / 15 minutes", "Always"),
    ("AW-Q4", "What is tomorrow’s first visible action?", "One line or Not ready to decide", "If unfinished task"),
    ("AW-Q5", "What are you still checking?", "Messages / Email / Tasks / News or social media / Something else", "If checking"),
    ("AW-Q6", "What does the next part of the evening need most?", "Quiet / Movement / Connection / Practical care / Nothing planned", "Before recovery choice"),
], [900, 3780, 3120, 1560], font_size=8.4)

heading(doc, "After Work routes", 2)
add_table(doc, ["Explicit state", "Primary action", "Optional instrument"], [
    ("Unfinished task", "Write tomorrow’s first visible action, place it where work resumes, then close the work surface.", "Short closure audio or relevant Observatory finding after completion."),
    ("Checking loop", "Choose a last-check time, move the device out of reach and begin one non-work action.", "Mirellis transition playlist after the device boundary."),
    ("Activated mind/body", "Use movement or a paced settling exercise matched to available time.", "Guided audio or non-screen sound environment."),
    ("Emotionally affected", "Name only what support is needed now—space, conversation, practical action or professional support. No interpretation.", "Optional user-authored note kept locally."),
    ("Depleted/numb", "Choose the lowest-demand recovery action: food, wash, rest, gentle movement or connection.", "Music may accompany; no pressure to reflect."),
    ("Cannot tell", "Neutral boundary: close one work cue, change physical context and choose one evening action.", "No diagnostic questions."),
], [2280, 4780, 2300], font_size=8.4)

heading(doc, "The immediate After Work sequence", 2)
number(doc, "Close one loop: save the return point or consciously choose not to.")
number(doc, "Create one physical boundary appropriate to location: close, pack, leave, change position or end the work channel.")
number(doc, "Begin one recovery action outside the app. If audio is selected, the audio supports the action rather than becoming a content destination.")
number(doc, "Return only to record: “Did work feel any more finished?” — Yes / A little / No / I am not sure.")

heading(doc, "Legacy 21-session program", 2)
paragraph(doc, "The existing program remains available under “After Work program.” Existing progress is migrated; it is not silently reset. The revised program is one session per practice day, not three unexplained exercises per day.")
add_table(doc, ["Phase", "Sessions", "Purpose", "Daily structure"], [
    ("Closure", "1–7", "Create a reliable return point and reduce unfinished-task residue.", "One return-point action + one boundary."),
    ("Boundary", "8–14", "Test physical, digital and social endings to work.", "One selected boundary + explicit outcome."),
    ("Recovery repertoire", "15–21", "Identify recovery actions the person deliberately wants to reuse.", "One action + optional save to My record."),
], [1500, 1100, 3590, 3170], font_size=8.8)
paragraph(doc, "Program language: “Session 1 of 21” and “one session today.” Never say “three exercises per day” if only one session is expected. Never use levels, matched decisions or adaptive difficulty without explaining the real consequence.", size=10, color=MUTED, italic=True)


heading(doc, "10. Night Audit and sleep-support experience", 1)
callout(doc, "Consumer job", "Help me understand what happens between intending to stop and actually stopping, then test one fitting change—without turning sleep into a score or treating every late night as procrastination.")

heading(doc, "Entry copy", 2)
paragraph(doc, "Headline: “You meant to stop. What happened next?”", bold=True)
paragraph(doc, "Support: “Reconstruct one recent night. Mirellis will suggest a possibility for you to confirm—not a diagnosis.”")

heading(doc, "Exact reconstruction questions", 2)
add_table(doc, ["Key", "Exact prompt", "Response control", "Requirement"], [
    ("NT-Q1", "What time did you intend to begin getting ready for sleep?", "Time picker / I had no intended time", "Required"),
    ("NT-Q2", "What time did you actually stop the activity that kept you up?", "Time picker / Approximate / I do not know", "Required if available"),
    ("NT-Q3", "What were you mainly doing in the gap?", "Scrolling or watching / Working / Talking or messaging / Gaming / Reading or making something / Household or caregiving / Lying awake / Something else", "Required"),
    ("NT-Q4", "At the time you intended to stop, what did your body feel like?", "Sleepy / Tired but mentally active / Not sleepy / Physically uncomfortable / Already in bed but awake / Not sure", "Required"),
    ("NT-Q5", "What did staying up give you in that moment?", "Time that felt mine / Relief or distraction / Connection / A chance to finish work / I continued without deciding / Nothing obvious / Not sure", "Required"),
    ("NT-Q6", "How often has this been happening recently?", "Less than weekly / 1–2 nights a week / 3 or more nights a week / Almost every night / Not sure", "Required"),
    ("NT-Q7", "How long has it been a concern?", "Under two weeks / 2–12 weeks / 3 months or longer / Not sure", "If 3+ nights weekly"),
], [900, 3780, 3220, 1460], font_size=8.2)

heading(doc, "Safety and suitability questions", 2)
paragraph(doc, "Use plain routing, not a scored clinical screen. The following appears only after the reconstruction and may be skipped.")
add_table(doc, ["Exact prompt", "Response", "System consequence"], [
    ("Has anyone noticed loud snoring, gasping, choking or pauses in your breathing during sleep?", "Yes / No / Not sure / Prefer not to say", "If Yes, explain that these signs deserve medical assessment; no apnea diagnosis."),
    ("Has sleepiness contributed to a near miss, accident or difficulty staying awake while driving?", "Yes / No / Prefer not to say", "If Yes, show immediate safety guidance and professional route."),
    ("Are sleep difficulties causing severe distress or making daily life difficult most days?", "Yes / No / Prefer not to say", "If Yes, prioritise professional support over a self-guided experiment."),
    ("Do you feel in immediate danger or unable to keep yourself safe?", "Yes / No / Prefer not to say", "If Yes, show locally appropriate emergency/crisis guidance; do not continue the audit."),
], [4740, 1900, 2720], font_size=8.4)

heading(doc, "Candidate patterns and routes", 2)
add_table(doc, ["Candidate possibility", "Explicit evidence required", "First experiment"], [
    ("Automatic continuation", "Continued without deciding + scrolling/watching/gaming + sleepy or tired.", "Live stopping cue, phone relocation and if–then plan."),
    ("Protected personal time", "Time that felt mine + delayed stopping.", "Schedule a protected decompression period earlier; keep bedtime decision separate."),
    ("Relief or avoidance", "Relief/distraction + mentally active or emotionally loaded.", "Short offload or calming action; professional route if distress is severe or persistent."),
    ("Social pull", "Connection + talking/messaging.", "Visible closing message and agreed stopping signal."),
    ("Unfinished work", "Working or chance to finish work.", "After Work return point before the Night action."),
    ("Timing mismatch", "Not sleepy at intended time + repeated pattern.", "Observe timing and consistency; do not force an earlier sleep window or label chronotype."),
    ("Possible sleep-onset difficulty", "Sleepy/tired, already in bed awake or persistent difficulty despite opportunity.", "General relaxation/stimulus information and clinician/CBT-I route when persistent."),
], [2550, 3860, 2950], font_size=8.3)

heading(doc, "Confirmation screen", 2)
paragraph(doc, "Exact heading: “One possibility—not a verdict.”", bold=True)
paragraph(doc, "Example body: “You said you were sleepy, kept scrolling and continued without making a new decision. One possibility is that the activity simply continued past the point you intended to stop.”")
paragraph(doc, "Controls: “This fits” / “This does not fit” / “Another possibility fits better” / “I do not want to label it.”")

heading(doc, "Night experiment protocols", 2)
add_table(doc, ["Protocol", "Structure", "What is compared", "Use"], [
    ("Quick test", "7 nights: nights 1–2 observe; nights 3–7 test one change.", "Intended stop, live/approximate actual stop and explicit ease/difficulty.", "Default free experiment."),
    ("Stronger comparison", "10 nights: 3-night observation + 7-night experiment.", "Same fields with a more stable personal baseline.", "Optional when the person wants clearer evidence."),
    ("No experiment", "One immediate action only.", "No before/after claim.", "When burden, fit or safety makes tracking inappropriate."),
], [1560, 3240, 3000, 1560], font_size=8.6)

heading(doc, "Live timestamp contract", 2)
bullet(doc, "The person may tap “I am stopping now” at the moment the activity ends. This creates a local timestamp.")
bullet(doc, "If completed the next morning, label the time “recalled” and never present it as measured.")
bullet(doc, "Do not ask the user to remain on-screen in bed. The action should reduce, not extend, device use.")
bullet(doc, "No wearable or health-platform integration is required at launch.")

heading(doc, "Night intervention boundaries", 2)
bullet(doc, "Sleep restriction, prescribed sleep windows, chronotherapy and medication/supplement guidance are excluded from self-guided V2.")
bullet(doc, "Sleep hygiene tips may support a broader plan but cannot be presented as treatment for chronic insomnia.")
bullet(doc, "Music, warm bathing, guided relaxation and to-do offload are optional instruments with qualified claims, not guaranteed sleep interventions.")
bullet(doc, "Persistent insomnia routes to CBT-I information and qualified practitioners.")


page_break(doc)
heading(doc, "11. Launch intervention library", 1)
paragraph(doc, "Internal keys support engineering and stewardship only. Consumer screens show the consumer name and approved explanation. Tier A means suitable for self-guided launch with comparatively stronger or directly relevant evidence; Tier B means optional/promising with qualified language; Tier C means information or clinician-guided only.")

heading(doc, "Tier definitions", 2)
add_table(doc, ["Tier", "Product treatment", "Claim form"], [
    ("A", "May be a primary self-guided recommendation when eligible.", "“This may make [specific action] easier because…”"),
    ("B", "Optional instrument or experiment; never the only route for a persistent problem.", "“Some people find… Evidence is limited or mixed.”"),
    ("C", "Do not deliver as a self-guided protocol. Explain and route to appropriate care.", "“A qualified professional can assess whether this is appropriate.”"),
    ("Excluded", "Not included in launch content.", "No consumer claim."),
], [900, 4750, 3710], font_size=8.8)

heading(doc, "Waking library", 2)
add_table(doc, ["Internal key", "Consumer action", "Tier", "Eligibility and approved purpose"], [
    ("WK-LIGHT", "Open light and movement", "B", "Recently awake, no safety flag and access to comfortable light/movement. Intended to support alertness; not a treatment."),
    ("WK-ORIENT", "Orient before accelerating", "B", "Anxious/rushed or uncertain first task. Intended to reduce immediate decision load."),
    ("WK-FIRST", "Begin one visible action", "A", "No clear start. Uses specific goal and immediate first action."),
    ("WK-CAPACITY", "Protect limited capacity", "A", "Short sleep opportunity or strong fatigue. Supports safer planning rather than promising recovery."),
    ("WK-CARE", "Look beyond a morning routine", "C", "Persistent, severe or safety-relevant sleepiness. Routes to assessment."),
], [1450, 2340, 750, 4820], font_size=8.7)

heading(doc, "Focus library", 2)
add_table(doc, ["Internal key", "Consumer action", "Tier", "Eligibility and approved purpose"], [
    ("FC-OUTCOME", "Define what must exist", "A", "Task unclear. Converts intention into an observable result."),
    ("FC-FIRST", "Name the first visible action", "A", "Start friction. Reduces ambiguity at the point of beginning."),
    ("FC-PHONE", "Put the phone beyond reach", "A", "Phone is not required or can be separated. Reduces a known distraction source."),
    ("FC-TOOLS", "Keep only the tools this task needs", "A", "Digital task with unnecessary tabs/apps. Environmental preparation."),
    ("FC-PARK", "Park the thought without following it", "B", "Internal task switching. Captures a brief user-authored cue for later."),
    ("FC-BOUNDARY", "Protect this work window", "A", "Interruption by people/messages. Creates an explicit availability boundary."),
    ("FC-TIMER", "Work for the time you actually have", "A", "Any session. User chooses duration; no universal 25- or 90-minute claim."),
    ("FC-MOVE", "Move before returning", "B", "Fatigue or prolonged sitting without safety concern. May support vigour; performance effects are modest/variable."),
    ("FC-MINDFUL", "Ten-minute attention reset", "B", "Optional for internal drift; some brief studies are promising, effects vary."),
    ("FC-STOP", "Stop or reschedule safely", "A", "Strong fatigue or safety concern. No discipline framing."),
    ("FC-GAME", "Cognitive training game", "Excluded", "Not included as a real-world focus solution; far-transfer claims remain uncertain."),
], [1450, 2440, 750, 4720], font_size=8.45)

heading(doc, "After Work and Restoration library", 2)
add_table(doc, ["Internal key", "Consumer action", "Tier", "Eligibility and approved purpose"], [
    ("AW-RETURN", "Leave tomorrow a starting point", "A", "Unfinished work. Records the next visible action and closes the current loop."),
    ("AW-CLOSE", "Close the work surface", "A", "Workspace available. Creates a physical boundary."),
    ("AW-DEVICE", "End the checking loop", "A", "Repeated checking. Last-check cue plus device relocation."),
    ("AW-MOVE", "Move out of work", "B", "Physical activation/restlessness. User-selected gentle movement."),
    ("AW-OFFLOAD", "Put the residue somewhere visible", "B", "Mental residue. Brief structured note; no interpretation."),
    ("AW-MUSIC", "Use sound as the boundary", "B", "User chooses music and can put the screen down. Instrument, not treatment."),
    ("AW-RELAX", "Follow a guided downshift", "B", "Activated mind/body. Relaxation or NSDR-style audio with restrained claims."),
    ("AW-RECOVER", "Choose what the evening needs", "A", "Depletion or unclear next state. Selects one practical recovery action."),
    ("AW-SUPPORT", "Ask for human support", "C", "Severe distress, persistent work harm or safety concern."),
], [1450, 2440, 750, 4720], font_size=8.45)

heading(doc, "Night library", 2)
add_table(doc, ["Internal key", "Consumer action", "Tier", "Eligibility and approved purpose"], [
    ("NT-LIVE", "Mark when you actually stop", "A", "Person wants a live timestamp. Local measurement with minimal interaction."),
    ("NT-IFTHEN", "Make the stopping decision in advance", "A", "Confirmed automatic continuation. One user-authored if–then plan."),
    ("NT-TIME", "Protect personal time earlier", "B", "Confirmed compensation/personal-time pattern. Tests whether earlier decompression changes the gap."),
    ("NT-PHONE", "Put the phone outside the stopping zone", "A", "Phone-driven continuation and device can be moved."),
    ("NT-WORK", "Close work before the night begins", "A", "Work residue. Routes through AW-RETURN."),
    ("NT-LIST", "Write tomorrow’s list, then stop planning", "B", "Unfinished tasks/worry. Based on limited laboratory evidence; no guaranteed latency claim."),
    ("NT-MUSIC", "Use a familiar sleep sound", "B", "User wants audio; subjective sleep benefits possible, objective effects uncertain."),
    ("NT-BATH", "Use warmth before bed", "B", "Safe access and no contraindication disclosed. Optional comfort routine with qualified evidence."),
    ("NT-RELAX", "Use a guided settling practice", "B", "Mentally active, no acute safety concern. Optional relaxation, not treatment."),
    ("NT-TIMING", "Observe when sleepiness actually arrives", "B", "Not sleepy at intended time. Observation only; no chronotherapy."),
    ("NT-CBTI", "Find evidence-based insomnia care", "C", "Persistent insomnia-like difficulty or significant impairment."),
    ("NT-APNEA", "Seek a sleep-breathing assessment", "C", "Gasping, choking, witnessed pauses or other concerning signs."),
], [1450, 2440, 750, 4720], font_size=8.35)

heading(doc, "Excluded launch content", 2)
for item in [
    "Medication, melatonin or supplement selection, dosing or timing.",
    "Sleep restriction, prescribed sleep windows and chronotherapy.",
    "Diagnostic ADHD, depression, burnout, chronotype, insomnia or apnea outputs.",
    "Guaranteed attention, sleep, nervous-system or brain-rewiring claims.",
    "White/pink noise superiority claims, binaural-beat treatment claims and proprietary-neuroscience claims without independent review.",
    "Universal hydration, posture, 25-minute or 90-minute productivity prescriptions.",
]:
    bullet(doc, item)


heading(doc, "12. Motivation, reminders and retention without manipulation", 1)
paragraph(doc, "Mirellis must help action happen, but “force” must mean deliberate precommitment and environmental friction—not coercion, guilt or punishment.")
add_table(doc, ["Mechanic to adopt", "How Mirellis uses it", "Boundary"], [
    ("One next action", "Today shows one relevant recommendation.", "User can always choose another or leave."),
    ("Implementation intention", "User confirms “If [cue], then I will [action].”", "Plan is user-authored or explicitly confirmed."),
    ("Appointment", "User selects a time and reminder mechanism.", "Explain reliability; no automatic default reminders."),
    ("Environmental commitment", "Phone placement, closed tabs, packed workspace or prepared light/movement context.", "Never claim enforcement that the platform cannot provide."),
    ("Visible completion", "Show the action and user-recorded outcome immediately.", "No score, grade or celebratory pressure."),
    ("Remembered context", "Tomorrow may show a deliberately saved practice or return point.", "Never silently advance a path."),
    ("Progressive challenge", "Only user-approved increases in session duration or complexity.", "No hidden level or algorithmic difficulty."),
], [2100, 4070, 3190], font_size=8.7)

heading(doc, "Reminder copy contract", 2)
add_table(doc, ["Use", "Approved example", "Prohibited example"], [
    ("Scheduled start", "“You asked Mirellis to remind you: begin the first visible action.”", "“Your streak is at risk.”"),
    ("Night stopping cue", "“You planned to stop now. Open Mirellis only if you want to mark the moment.”", "“It is bedtime. Put your phone down now.”"),
    ("After Work", "“You saved a return point. Close work when you are ready.”", "“Complete today’s recovery session.”"),
    ("Missed reminder", "No follow-up unless the user requested one.", "“You missed yesterday.”"),
], [1860, 4050, 3450], font_size=8.8)

heading(doc, "Retention definition", 2)
paragraph(doc, "Retention is not “open the app every day.” Product success is returning when a relevant boundary becomes difficult, reusing a saved action without dependence, or no longer needing Mirellis for a learned practice. The app may end an experiment and encourage independence.")


heading(doc, "13. Personal Observatory integration", 1)
callout(doc, "Role", "My record is a private mirror, not a performance dashboard. It contains only what the person completed, deliberately saved or explicitly confirmed.")

heading(doc, "True empty state", 2)
paragraph(doc, "Heading: “Nothing yet.”", bold=True)
paragraph(doc, "Body: “Complete or deliberately save one practice to begin a private record you can use without Mirellis.”")
paragraph(doc, "Primary action: “Back to today.”")

heading(doc, "Allowed record types", 2)
add_table(doc, ["Record", "Created when", "Displayed evidence"], [
    ("Completed action", "User completes and confirms the close.", "Action name, date, explicit outcome and optional note."),
    ("Saved practice", "User taps a save control after seeing its consequence.", "Practice name, when to use it and saved reminder if any."),
    ("Confirmed pattern", "User chooses “This fits.”", "The exact answers that support it and the user’s confirmation date."),
    ("Experiment comparison", "Required observations exist and user opens the comparison.", "Raw nights/sessions and plain differences; no score or verdict."),
    ("Personal learning", "User writes or selects a statement and confirms it.", "User’s words plus linked local evidence."),
], [1850, 3440, 4070], font_size=8.8)

heading(doc, "Personal-learning lifecycle", 2)
number(doc, "Candidate: the system offers a possibility linked to explicit evidence.")
number(doc, "Confirmed: the person selects “This fits.”")
number(doc, "Kept: the person chooses to retain it in My record.")
number(doc, "Questioned: the person marks it uncertain or adds their own note.")
number(doc, "Revised: the person authors a replacement statement; the old version remains visible until deletion if they choose.")
number(doc, "Discarded: removed locally and from optional sync according to the deletion contract.")

heading(doc, "Comparison language", 2)
bullet(doc, "Allowed: “On the first two nights, you recalled stopping 42 and 55 minutes after your intended time. On four of the five experiment nights, you marked the stop within 25 minutes.”")
bullet(doc, "Not allowed: “Your bedtime discipline improved by 62%.”")
bullet(doc, "Allowed: “In two sessions with the phone in another room, you selected ‘no phone interruption.’”")
bullet(doc, "Not allowed: “Your focus score rose.”")


heading(doc, "14. Practitioner directory and safety routing", 1)
paragraph(doc, "The directory is an accountable route out of the app when self-guided action is insufficient or inappropriate. It is not a monetised referral disguised as clinical advice.")
add_table(doc, ["Trigger", "Consumer message", "Directory filters"], [
    ("Persistent sleep-onset or maintenance difficulty", "“Persistent sleep difficulty can have several causes. Evidence-based insomnia care may be a better next step than another self-guided experiment.”", "Sleep medicine; CBT-I; telehealth/in-person; language; location."),
    ("Gasping, choking or witnessed breathing pauses", "“These signs deserve medical assessment. Mirellis cannot determine their cause.”", "Sleep medicine; pulmonology; diagnostic services."),
    ("Dangerous daytime sleepiness", "“A quick routine cannot make severe sleepiness safe. Please choose a safer immediate plan and seek assessment.”", "Sleep medicine; primary care; urgent local guidance where appropriate."),
    ("Persistent attention difficulty across settings", "“Focus difficulty can have many causes. A qualified professional can assess what is relevant in your case.”", "Clinical psychology; psychiatry; occupational health; primary care."),
    ("Severe distress or immediate safety concern", "“Mirellis is not emergency support.”", "Local emergency/crisis resources; no commercial directory interstitial."),
], [2170, 4880, 2310], font_size=8.4)

heading(doc, "Directory data boundaries", 2)
bullet(doc, "No audit answers, notes, pattern labels or experiment results are placed in referral URLs.")
bullet(doc, "The person may export a private summary and choose to share it themselves.")
bullet(doc, "Paid listings, if introduced, must be visibly labelled and must not change safety-route ordering.")
bullet(doc, "Verification status, profession, location and scope must be factual and maintained separately from the app’s intervention engine.")


heading(doc, "15. Data model, privacy, consent and deletion", 1)
heading(doc, "Recommended launch decision", 2)
callout(doc, "Local-first", "No account is required to begin or complete any launch experience. IndexedDB is the primary store. Optional encrypted cloud backup and cross-device recovery are deferred until the person deliberately enables an account.")

heading(doc, "Core objects", 2)
add_table(doc, ["Object", "Minimum fields", "Prohibited derived data"], [
    ("profile_preferences", "locale, timezone, accessibility preferences, reminder permissions", "Personality, readiness, chronotype or health risk score"),
    ("practice_session", "session_id, domain, practice_key, content_version, started_at, completed_at/left_at, explicit outcome", "Quality score or inferred success"),
    ("answer", "question_key, option_key or user text, authored_at, source=explicit", "Sentiment, diagnosis or topic inference"),
    ("candidate_pattern", "pattern_key, evidence answer keys, presented_at, confirmation state", "Unconfirmed personal label"),
    ("experiment", "protocol_key, selected intervention, start/end, observations, status", "Clinical efficacy conclusion"),
    ("observation", "user-authored/selected statement, linked local evidence, kept/discarded state", "Automatic insight"),
    ("reminder", "mechanism, schedule, permission status, last delivery result", "Penalty or compliance score"),
    ("consent", "purpose, version, scope, granted/revoked time", "Bundled or implied consent"),
], [1700, 4990, 2670], font_size=8.35)

heading(doc, "Local retention defaults", 2)
bullet(doc, "Unfinished sessions: retained locally for 30 days unless completed, deleted or resumed.")
bullet(doc, "Completed sessions and saved practices: retained until the person deletes them.")
bullet(doc, "Unsaved free text: discarded when the person leaves the entry flow, after a clear warning if text would be lost.")
bullet(doc, "Safety responses: used for immediate routing; not retained unless the person deliberately saves the session summary.")
bullet(doc, "No private data is added to logs, analytics events, error traces, URLs or notification bodies.")

heading(doc, "Export", 2)
paragraph(doc, "Offer human-readable PDF/HTML and machine-readable JSON. Export must identify recalled versus live timestamps, content versions, user-authored notes and user-confirmed patterns. It must not add interpretations.")

heading(doc, "Deletion", 2)
number(doc, "User selects “Delete my local Mirellis data.”")
number(doc, "Show a cancelable confirmation naming exactly what will be removed and what, if anything, exists in optional cloud backup.")
number(doc, "Primary destructive action: “Delete data.” Secondary action: “Cancel.” Initial focus remains on Cancel.")
number(doc, "After confirmation, delete local objects, caches containing personal state and optional synced copies. Preserve only legally required consent/audit records, if applicable, with explanation.")
number(doc, "Announce completion through visible and aria-live text: “Your Mirellis data was deleted from this device.”")


heading(doc, "16. AI boundaries and content governance", 1)
paragraph(doc, "AI is not required for the first V2 release. The governed library and deterministic rules are sufficient for meaningful personalisation and are easier to test, explain and audit.")
add_table(doc, ["AI may later…", "AI must never…"], [
    ("Rewrite a user-authored first action for clarity and ask for confirmation.", "Diagnose a sleep, attention or mental-health condition."),
    ("Explain why an already-selected, approved intervention fits explicit answers.", "Select an intervention outside deterministic eligibility rules."),
    ("Offer two approved copy variants matched to preferred tone.", "Interpret diary notes, infer emotion or create hidden user attributes."),
    ("Summarise an export using only user-confirmed records, with confirmation.", "Generate medical advice, supplement guidance or research findings."),
], [4680, 4680], font_size=8.8)

heading(doc, "Content release workflow", 2)
number(doc, "Product owner writes the user problem, action and measurable completion evidence.")
number(doc, "Research owner assigns evidence tier, approved claim and source register entries.")
number(doc, "Clinical reviewer approves safety wording and exclusions where the content touches sleep or mental health.")
number(doc, "Trust reviewer confirms privacy, consent and non-manipulation compliance.")
number(doc, "Engineering implements the versioned object and automated branch tests.")
number(doc, "QA verifies every consumer control, copy state and accessibility state before promotion.")

heading(doc, "Claim governance", 2)
bullet(doc, "Every claim object has exact public copy, evidence tier, source, review date and expiry/review date.")
bullet(doc, "A weaker evidence tier cannot be upgraded by marketing language.")
bullet(doc, "User testimonials cannot be used as clinical evidence.")
bullet(doc, "Transcript statements remain hypotheses until linked to primary or authoritative evidence.")


heading(doc, "17. Engineering architecture, PWA/native boundaries and cost", 1)
heading(doc, "Architecture recommendation", 2)
add_table(doc, ["Layer", "Launch recommendation", "Reason"], [
    ("Frontend", "Preserve the existing app framework if sound; otherwise use a typed React/Next.js application after repository audit.", "Avoid needless rewrite while enabling reusable flows and PWA support."),
    ("Content", "Versioned TypeScript/JSON registry in the repository for V2 launch.", "Reviewable diffs, deterministic deployment and no CMS dependency."),
    ("Rules", "Pure deterministic functions with exhaustive fixtures.", "Explainable routing and testability."),
    ("Local data", "IndexedDB with a versioned migration layer; minimal localStorage only for non-sensitive flags.", "Multi-session persistence and reliable structured storage."),
    ("Cloud", "None required for launch; optional Supabase Auth/Postgres/Storage later with RLS.", "Protects no-account entry and keeps launch cost/complexity low."),
    ("Email", "None required for core action; Resend only for optional recovery/account links after opt-in.", "Avoids email dependency and unsolicited reminders."),
    ("Hosting", "Vercel preview and production environments with spend controls.", "Matches current deployment and provides rollback."),
    ("AI", "No launch dependency. Later, server-side gateway with structured outputs and no note retention.", "Prevents unpredictable medical/personal interpretation."),
], [1550, 4800, 3010], font_size=8.3)

heading(doc, "PWA capability matrix", 2)
add_table(doc, ["Capability", "Web/PWA launch", "Native later"], [
    ("Offline practices and records", "Yes, after first load via service worker and IndexedDB.", "Yes"),
    ("Installable home-screen app", "Yes where browser/platform supports installation.", "App-store installation"),
    ("Reliable scheduled notifications", "Conditional; platform/browser restrictions must be explained.", "More reliable platform notifications"),
    ("Block other apps", "No", "Possible through approved OS APIs and entitlements"),
    ("Background timers", "Limited and may be throttled.", "Native background mechanisms subject to OS policy"),
    ("Widgets/lock-screen actions", "No or highly limited", "Possible"),
    ("Health/wearable data", "Not included", "Separate consented future capability"),
], [2580, 3620, 3160], font_size=8.7)

heading(doc, "Cost envelope as of August 2026", 2)
add_table(doc, ["Stage", "Expected base services", "Indicative monthly platform cost"], [
    ("Local-first preview", "Vercel preview; repository content; no cloud accounts.", "$0 during controlled noncommercial preview if plan terms permit."),
    ("Commercial production, local-first", "Vercel Pro.", "From $20/month plus usage and tax."),
    ("Optional cloud sync", "Vercel Pro + Supabase Pro.", "From approximately $45/month plus usage and tax."),
    ("Optional email at low volume", "Resend Free within published limits.", "$0; current free allowance 3,000 emails/month and 100/day."),
    ("Email after free allowance", "Resend Pro.", "From $20/month for 50,000 emails."),
], [2390, 4090, 2880], font_size=8.7)
paragraph(doc, "Pricing changes. Recheck official pricing before production approval. Vercel Hobby is for personal/non-commercial use; a commercial Mirellis launch should budget for Pro.", size=9.5, color=MUTED, italic=True)

heading(doc, "Security requirements", 2)
bullet(doc, "Content Security Policy, secure headers, dependency scanning and secret isolation.")
bullet(doc, "No personal answers in client/server logs, error trackers or analytics payloads.")
bullet(doc, "If cloud sync launches: row-level security, per-user ownership tests, encrypted transport, deletion jobs and backup policy disclosure.")
bullet(doc, "No AI prompt may contain private notes unless the person deliberately invokes a future AI feature and confirms transmission.")
bullet(doc, "Deployment rollback must preserve local schema compatibility.")


heading(doc, "18. Measurement and learning operations", 1)
paragraph(doc, "Measurement exists to determine whether the product functions and helps people complete intended actions—not to maximise time in app. The current build must continue to state “Analytics are not active in this build” until measurement is actually implemented and disclosed.")

heading(doc, "Recommended launch measurement", 2)
callout(doc, "Consent decision", "Use opt-in, data-minimised product measurement. Never collect notes, answer text, safety responses, directory searches or private pattern labels in analytics.")
add_table(doc, ["Event", "Permitted properties", "Purpose"], [
    ("app_opened", "app version; anonymous install ID if consented; platform class", "Reliability and aggregate return."),
    ("moment_selected", "moment key only", "Entry distribution."),
    ("recommendation_shown", "intervention key; content version", "Routing coverage and missing routes."),
    ("practice_started", "intervention key; planned duration", "Start friction."),
    ("practice_completed", "intervention key; elapsed class", "Completion and technical failure."),
    ("practice_left", "intervention key; screen key", "Identify confusing or burdensome steps."),
    ("returned_later", "days since prior session; no personal outcome", "Day-return behavior."),
    ("permission_result", "permission type; granted/denied/dismissed", "Permission UX quality."),
    ("error_occurred", "error code; screen; app version", "Reliability."),
], [2200, 4300, 2860], font_size=8.65)

heading(doc, "Forbidden analytics", 2)
for item in [
    "Raw or hashed free text, task titles, notes or tomorrow return points.",
    "Exact sleep/wake timestamps, health/safety answers or inferred patterns.",
    "Session replay on private input screens.",
    "Advertising identifiers, cross-site tracking or sale/share of personal data.",
    "Engagement experiments that add guilt, streaks or pressure.",
]:
    bullet(doc, item)

heading(doc, "Learning review cadence", 2)
bullet(doc, "Weekly during preview: defects, dead ends, confusing copy and failed permissions.")
bullet(doc, "After the first 30 complete users per moment: qualitative comprehension review and branch coverage audit.")
bullet(doc, "After sufficient experiment completions: review retention and usefulness without claiming population-level research.")
bullet(doc, "Any public Finding requires a separate research protocol, consent basis and methodology review; product data alone is not an Observatory research finding.")


page_break(doc)
heading(doc, "19. Build order and release plan", 1)
paragraph(doc, "Work in small, reviewable commits. Each phase ends with a preview deployment, copy/state logic report and automated plus manual evidence. No phase is promoted solely because it builds successfully.")
add_table(doc, ["Phase", "Scope", "Exit gate"], [
    ("0. Repository and data audit", "Document current routes, framework, storage schema, live V1 state, deployment pipeline and all existing user data migrations.", "Canonical architecture report; backup/rollback plan; no destructive migration."),
    ("1. Shared foundation", "New app shell, Today/Practices/My record/Help/Settings, design tokens, state machine, content registry, local database and migrations.", "Mobile/keyboard/offline shell passes; no consumer internal IDs."),
    ("2. V1 After Work retrofit", "Problem-first entry, truthful routes, one-session language, legacy 21-session migration and clear reminder mechanism.", "Every branch works; old progress preserved; no matched-decision language."),
    ("3. Night Audit", "Reconstruction, safety route, pattern confirmation, 7- and 10-night experiments, live/recalled timestamps.", "Clinical wording reviewed; all safety branches pass; no diagnosis."),
    ("4. Focus", "Real-result session, preparation routes, active timer, interruption capture and 3-session experiment.", "Users can explain the task and begin real work; platform blocking claims truthful."),
    ("5. Waking", "Immediate action routes, capacity/safety route and optional 7-morning experiment.", "No universal morning cure claim; safety-critical route tested."),
    ("6. Observatory and directory", "True empty state, explicit saves, evidence-linked records, export/delete and privacy-preserving directory links.", "Record provenance, cancel/confirm deletion and export verified."),
    ("7. PWA hardening", "Install, offline cache, update path, notification explanation, recovery from interrupted updates.", "Cross-platform capability matrix tested on real devices."),
    ("8. Controlled release", "Feature flags, consented measurement and invited-user preview.", "Acceptance criteria pass; owner approves production promotion."),
], [1750, 5000, 2610], font_size=8.2)

heading(doc, "Required commit report", 2)
bullet(doc, "What changed, including files and affected routes.")
bullet(doc, "Before/after consumer copy.")
bullet(doc, "State logic and data migration changed.")
bullet(doc, "Automated test evidence and manual device/accessibility evidence.")
bullet(doc, "Known limitations and whether deployment is preview-only or production-approved.")

heading(doc, "Migration principles", 2)
bullet(doc, "Never delete or reset an existing practice record to simplify the new schema.")
bullet(doc, "Map legacy V1 sessions to the new practice_session object with a source_version field.")
bullet(doc, "If a legacy field has no truthful equivalent, preserve it in a legacy payload accessible to export, not in current consumer summaries.")
bullet(doc, "Migrations are idempotent and tested against empty, first-time, partial, completed and corrupted local states.")


heading(doc, "20. Accessibility and inclusive design", 1)
add_table(doc, ["Area", "Requirement"], [
    ("Targets", "Minimum 44 × 44 CSS px for every interactive control; adequate separation."),
    ("Keyboard", "All actions available by keyboard; logical focus order; no keyboard trap; Escape only when behavior is explained."),
    ("Focus", "Visible focus indicator with at least 3:1 contrast against adjacent colors."),
    ("Semantics", "Native buttons, inputs, radio groups, headings, landmarks and status regions."),
    ("Status", "aria-live polite for save/completion/reminder status; assertive only for immediate safety/error messages."),
    ("Motion", "Respect prefers-reduced-motion; timers and state changes remain understandable without animation."),
    ("Color", "Never rely on color alone; text/background and control contrast meet WCAG AA."),
    ("Time", "Timers can be paused/extended where the action permits; no forced timeout for reading or answering."),
    ("Language", "Plain English; no arrows, abbreviations, clinical jargon or unexplained metaphors."),
    ("Forms", "One conceptual decision per screen where possible; errors identify the issue and preserve entered data."),
    ("Audio", "Transcript or equivalent instructions; playback controls labelled; no autoplay."),
    ("Mobile", "320 px width without horizontal scrolling; safe-area support; zoom to 200% without loss."),
], [1650, 7710], font_size=8.6)

heading(doc, "Cognitive accessibility", 2)
bullet(doc, "Explain why each question is asked when it is not self-evident.")
bullet(doc, "Show exact session duration before beginning.")
bullet(doc, "Avoid simultaneous choices about problem, intervention, reminder and saving.")
bullet(doc, "Replace abstract classifications with familiar descriptions and examples.")
bullet(doc, "Allow “Not sure” without forcing a false answer or producing a generic verdict disguised as personalisation.")


heading(doc, "21. Test plan and acceptance criteria", 1)
heading(doc, "Automated tests", 2)
add_table(doc, ["Test layer", "Required coverage"], [
    ("Unit", "Every deterministic routing rule, exclusion, recommendation reason, progress label and completion consequence."),
    ("Content schema", "Every intervention has required fields, valid evidence tier, approved claim and no forbidden consumer identifier."),
    ("State machine", "Start, pause, leave, resume, complete, save, unsave, edit, discard and expire."),
    ("Migration", "All legacy V1 local states, repeated migration and rollback-compatible reads."),
    ("Privacy", "Analytics payload allowlist; no answers/notes/timestamps in logs, URLs or events."),
    ("Accessibility", "Automated axe-style checks plus target-size and reduced-motion assertions."),
    ("Offline/PWA", "First load, cached route, offline action, reconnect, update and schema upgrade."),
    ("Security", "Dependency, headers, injection, XSS and storage-boundary checks."),
], [1850, 7510], font_size=8.8)

page_break(doc)
heading(doc, "Manual end-to-end matrix", 2)
add_table(doc, ["Scenario", "Required evidence"], [
    ("First-time user", "Understands purpose, begins one action without account and can leave."),
    ("All Waking routes", "Each answer produces the specified action or safety boundary."),
    ("All Focus routes", "Real result, clear preparation, timer, close and no score."),
    ("All After Work routes", "Materially different actions; legacy program/progress still accessible."),
    ("All Night routes", "Pattern confirmation, nonfit route, safety exits, live/recalled timestamp distinction."),
    ("Interrupted session", "State preserved; next label says Continue; no silent completion."),
    ("Saved practice", "Immediate visible result in My record and truthful Tomorrow consequence."),
    ("Nothing saved", "True “Nothing yet” empty state; no fake content."),
    ("Reflection", "Optional note saves, edits, exports and deletes without interpretation."),
    ("Reminder", "Mechanism explained; permission result handled; edit/disable works."),
    ("Deletion cancel", "No data removed; focus returns safely."),
    ("Deletion confirm", "All scoped data removed and status announced."),
    ("Keyboard only", "Logical order and activation across every real branch."),
    ("Screen reader", "Prompts, groups, timer, save and completion feedback announced meaningfully."),
    ("Reduced motion", "No essential information lost and no disorienting transition."),
    ("Mobile devices", "iOS Safari/PWA, Android Chrome/PWA and 320 px responsive state."),
], [2380, 6980], font_size=8.55)

heading(doc, "Product acceptance criteria", 2)
for item in [
    "A new user can state what Mirellis is for after the first screen.",
    "A user reaches a real action in no more than four required decisions for an immediate session.",
    "Every visible answer changes routing, explanation, eligibility or stored explicit evidence.",
    "Every visible control works and every label describes its actual consequence.",
    "No user is told to tap text that is not a control.",
    "No untouched practice is labelled Continue and no completed practice is silently advanced.",
    "No consumer screen exposes internal IDs, levels, matched decisions, transition arrows or curriculum terms.",
    "A person can complete every core experience without account, purchase, research consent or analytics consent.",
    "Every clinical/safety concern exits to accurate help without a diagnosis.",
    "All personal insights are authored, selected or confirmed and linked to explicit evidence.",
    "The live deployment can be rolled back without losing or corrupting existing local V1 progress.",
]:
    bullet(doc, item)

heading(doc, "Usability acceptance", 2)
bullet(doc, "Five representative first-time users per moment can explain what the recommended action requires before starting it.")
bullet(doc, "At least four of five can complete or deliberately leave the core flow without moderator explanation.")
bullet(doc, "No participant confuses a reminder with a guaranteed alarm or misunderstands sessions/days/program duration.")
bullet(doc, "Confusion is treated as a product defect, not user failure; wording and flow are revised before wider release.")


heading(doc, "22. Approval decisions and version freeze", 1)
heading(doc, "Recommended decisions for approval", 2)
add_table(doc, ["Decision", "Recommended V1 + V2 position", "Consequence"], [
    ("Account", "No account required; optional cloud recovery deferred.", "Launch remains local-first and private."),
    ("Mood screening", "Do not implement PHQ-2 or diagnostic scoring in launch.", "Use direct safety/support prompts and clinically review future screening."),
    ("Analytics", "Opt-in minimal telemetry only after disclosure and consent.", "Current Settings keeps “Analytics are not active” until implemented."),
    ("V1 program", "Preserve and migrate the 21-session After Work program; make it optional maintenance.", "No existing progress loss; one session/day language."),
    ("Subscription", "None at launch.", "All core features are free while usefulness and completion are tested."),
    ("AI", "No launch dependency.", "Deterministic routing ships first."),
    ("Native app", "PWA first; native only after use justifies OS enforcement/reminders.", "No false app-blocking claims."),
    ("Clinical review", "Required for Night safety/routing and any future validated screen.", "No production Night release before sign-off."),
], [2000, 4530, 2830], font_size=8.4)

heading(doc, "Version freeze", 2)
callout(doc, "Freeze", "Version 1.0 contains four problem-first consumer moments, one shared action engine, the preserved After Work program, Night Audit, Waking and Focus experiences, Personal Observatory records and a professional-help route. Do not add extra curriculum, subscriptions, scores, streaks, social features, AI interpretation, supplements, wearable inference or native enforcement during this build.")

heading(doc, "What must be approved before code", 2)
number(doc, "The four consumer moments and the status of Night Audit as a V2 capability rather than a fifth V1 transition.")
number(doc, "The exact question and answer registries in Sections 7–10.")
number(doc, "The launch intervention library and evidence tiers in Section 11.")
number(doc, "The recommended local-first/no-account architecture and optional telemetry position.")
number(doc, "The build order and acceptance gates.")
number(doc, "The requirement for clinical review before the Night experience reaches production.")


page_break(doc)
heading(doc, "Appendix A. Event and object dictionaries", 1)
heading(doc, "Practice state values", 2)
add_table(doc, ["Value", "Meaning", "Allowed next states"], [
    ("untouched", "Content has never been started.", "active, left"),
    ("active", "Current step is in progress.", "paused, left, completed"),
    ("paused", "User deliberately paused and state is resumable.", "active, left, completed"),
    ("left", "User left without completion.", "active, expired, discarded"),
    ("completed", "Required action close was recorded.", "saved, discarded, returned"),
    ("saved", "User deliberately kept the practice/record.", "unsaved, returned, discarded"),
    ("expired", "Unfinished local retention period ended.", "discarded"),
    ("discarded", "Removed from active/personal record scope.", "none"),
], [1600, 4580, 3180], font_size=8.8)

heading(doc, "Pattern confirmation values", 2)
add_table(doc, ["Value", "Consumer meaning", "Record consequence"], [
    ("presented", "Possibility shown but not answered.", "Not part of My record."),
    ("confirmed", "“This fits.”", "Eligible to be saved with linked explicit evidence."),
    ("rejected", "“This does not fit.”", "Not shown as personal learning."),
    ("alternative", "Another governed possibility selected.", "Only selected possibility proceeds."),
    ("unlabelled", "“I do not want to label it.”", "No pattern stored; action may still be chosen."),
], [1700, 3590, 4070], font_size=8.8)

heading(doc, "Minimum export schema", 2)
add_table(doc, ["Collection", "Fields"], [
    ("metadata", "exported_at, app_version, schema_version, timezone, scope"),
    ("sessions", "session_id, moment, practice name/key, content version, explicit times, state, explicit outcome"),
    ("answers", "question key/label, selected option or authored text, authored_at, live/recalled provenance"),
    ("patterns", "candidate label, evidence references, confirmation state/date"),
    ("experiments", "protocol, observations, selected change, comparison presentation"),
    ("observations", "user-authored/selected statement, linked evidence, kept/discarded state"),
    ("consents", "purpose, policy version, granted/revoked time and scope"),
], [1850, 7510], font_size=8.8)


heading(doc, "Appendix B. Evidence register", 1)
paragraph(doc, "This register supports product design; it is not a substitute for clinical review. Product claims must use the narrower interpretation recorded in this specification.")
source_line(doc, 1, "AASM: Behavioral and psychological treatments for chronic insomnia", "https://aasm.org/new-guideline-supports-behavioral-psychological-treatments-for-insomnia/", "CBT-I is the strong recommendation; sleep hygiene is not a standalone treatment.")
source_line(doc, 2, "Bedtime procrastination intervention RCT (2023)", "https://pubmed.ncbi.nlm.nih.gov/37354745/", "Promising reduction in a small, nonclinical young-adult sample; supports a bounded experiment, not a universal sleep product.")
source_line(doc, 3, "Bedtime procrastination systematic review and meta-analysis", "https://pubmed.ncbi.nlm.nih.gov/36375334/", "Associations are plausible but overall evidence quality is low and causal certainty is limited.")
source_line(doc, 4, "Mental contrasting with implementation intentions for bedtime procrastination", "https://pubmed.ncbi.nlm.nih.gov/31403339/", "Supports user-authored obstacle and if–then planning for this specific behavior.")
source_line(doc, 5, "Daily versus one-time MCII registered report", "https://pubmed.ncbi.nlm.nih.gov/40219795/", "Supports repeated daily application rather than assuming one exposure is enough.")
source_line(doc, 6, "Smartphone brain-drain meta-analysis", "https://pubmed.ncbi.nlm.nih.gov/37754029/", "Overall negative effect of smartphone presence/use with domain and context heterogeneity.")
source_line(doc, 7, "Acute exercise and cognition meta-review", "https://pubmed.ncbi.nlm.nih.gov/39883421/", "Small-to-medium acute effects; supports optional movement with modest claims.")
source_line(doc, 8, "Psychological detachment intervention meta-analysis", "https://pubmed.ncbi.nlm.nih.gov/34096763/", "Modest average benefit; supports After Work as a real restoration domain.")
source_line(doc, 9, "Micro-breaks meta-analysis", "https://pubmed.ncbi.nlm.nih.gov/36044424/", "Small benefits for vigour/fatigue; no reliable overall demanding-task performance benefit.")
source_line(doc, 10, "Sleep timing and consistency systematic review", "https://pubmed.ncbi.nlm.nih.gov/33054339/", "Regularity is favorably associated with health, but thresholds and much causal evidence remain limited.")
source_line(doc, 11, "Morning light for delayed sleep-wake phase disorder review", "https://pubmed.ncbi.nlm.nih.gov/34381579/", "Protocols and efficacy are not settled; avoid universal light prescriptions or circadian treatment claims.")
source_line(doc, 12, "Reactive countermeasures to sleep inertia review", "https://pubmed.ncbi.nlm.nih.gov/27193071/", "No convincing universal reactive countermeasure; safety-critical tasks should not rely on a quick routine.")
source_line(doc, 13, "Warm shower or bath before sleep meta-analysis", "https://pubmed.ncbi.nlm.nih.gov/31102877/", "Potential benefit under studied timing/temperature conditions; limited protocol certainty.")
source_line(doc, 14, "Bedtime to-do list laboratory study", "https://pubmed.ncbi.nlm.nih.gov/29058942/", "Specific promising finding; insufficient for guaranteed or population-wide claims.")
source_line(doc, 15, "Music for insomnia Cochrane review", "https://pubmed.ncbi.nlm.nih.gov/26270746/", "May improve subjective sleep quality; evidence for objective sleep outcomes is limited.")
source_line(doc, 16, "Yoga Nidra systematic review", "https://pubmed.ncbi.nlm.nih.gov/41144325/", "Promising but heterogeneous, small and at moderate/high risk of bias with weak safety reporting.")
source_line(doc, 17, "Brief daily meditation RCT", "https://pubmed.ncbi.nlm.nih.gov/30153464/", "Thirteen minutes daily improved selected outcomes after eight weeks, not four; does not support instant permanent rewiring.")
source_line(doc, 18, "Attentional-blink intensive meditation study", "https://pmc.ncbi.nlm.nih.gov/articles/PMC1865565/", "The influential study involved 17 participants and three months of intensive practice, not a permanent 17-minute intervention.")
source_line(doc, 19, "Goal-setting systematic review and meta-analysis", "https://pubmed.ncbi.nlm.nih.gov/29189034/", "Small positive unique effect across behaviors; supports specific, observable goals with context-sensitive use.")
source_line(doc, 20, "Commercial brain-training systematic review", "https://pubmed.ncbi.nlm.nih.gov/34251578/", "Near-transfer may occur; evidence for broad objective everyday benefit is insufficient.")
source_line(doc, 21, "FTC action concerning Lumosity claims", "https://search.ftc.gov/news-events/news/press-releases/2016/01/lumosity-pay-2-million-settle-ftc-deceptive-advertising-charges-its-brain-training-program", "Regulatory warning against unsupported cognitive and health claims.")
source_line(doc, 22, "AASM position on home sleep apnea testing", "https://aasm.org/advocacy/position-statements/clinical-use-of-a-home-sleep-apnea-test-an-updated-american-academy-of-sleep-medicine-position-statement/", "Medical providers diagnose OSA; automated data alone must not drive diagnosis/treatment.")
source_line(doc, 23, "USPSTF statement on OSA screening", "https://www.uspreventiveservicestaskforce.org/uspstf/document/RecommendationStatementFinal/obstructive-sleep-apnea-in-adults-screening", "Insufficient evidence for screening asymptomatic adults; symptomatic concerns require appropriate clinical management.")
source_line(doc, 24, "Opal explanation of app blocking", "https://opalapp.com/help/what-is-opal", "iOS blocking depends on Apple Screen Time APIs; informs the PWA/native boundary.")
source_line(doc, 25, "Strava Instant Workouts", "https://support.strava.com/en-us/articles/15401583-instant-workouts", "Useful product pattern: explicit focus, personalized action, estimated effort, purpose and immediate recording.")
source_line(doc, 26, "Vercel pricing", "https://vercel.com/pricing", "Pro starts at $20/month; Hobby is positioned for personal/noncommercial use.")
source_line(doc, 27, "Supabase pricing", "https://supabase.com/pricing", "Free tier exists; Pro starts at $25/month and free projects may pause after inactivity.")
source_line(doc, 28, "Resend pricing", "https://resend.com/docs/knowledge-base/what-is-resend-pricing", "Free tier currently includes 3,000 emails/month and 100/day; Pro begins at $20/month.")

heading(doc, "Evidence interpretation rules", 2)
bullet(doc, "A single laboratory result can support an optional test, not a broad promise.")
bullet(doc, "Associational evidence can support observation or a question, not causal explanation.")
bullet(doc, "A clinical intervention cannot be converted into a self-guided micro-practice merely because one component appears simple.")
bullet(doc, "Product telemetry is not research evidence and cannot become an Observatory Finding without a separate protocol.")

heading(doc, "Document approval block", 1)
add_table(doc, ["Role", "Name", "Decision / date"], [
    ("Founder / product owner", "", "Approve / Amend / Reject"),
    ("Clinical reviewer — sleep", "", "Approve / Amend / Reject"),
    ("Trust and privacy reviewer", "", "Approve / Amend / Reject"),
    ("Engineering owner", "", "Feasible / Risks noted"),
], [2600, 2900, 3860], font_size=9.3)


OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
