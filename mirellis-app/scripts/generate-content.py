"""Generate the configuration-driven transition content from frozen product specs."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/aseem/Documents/Mirellis App')
DELIVERABLES = Path('/Users/aseem/.codex/.chatgpt-projects/g-p-68a583386d9481918840fb251b9a3ede/deliverables')

SPECS = [
    {
        'file': 'Transition 001 – Sleep → Wakefulness Product Specification v1.0.docx',
        'id': 'sleep-wakefulness',
        'title': 'Sleep → Wakefulness',
        'question': 'Can I make a workable start after waking?',
        'prefix': 'SW',
        'first_table': 11,
        'stride': 2,
        'states': ['Orient', 'Choose', 'Influence', 'Activate', 'Self-direct'],
    },
    {
        'file': 'Transition 002 – Wakefulness → Readiness Product Specification v1.0.docx',
        'id': 'wakefulness-readiness',
        'title': 'Wakefulness → Readiness',
        'question': 'Can I choose how the day begins?',
        'prefix': 'WR',
        'first_table': 10,
        'stride': 2,
        'states': ['Arrive', 'Protect', 'Orient', 'Begin', 'Self-direct'],
    },
    {
        'file': 'Transition 003 – Readiness → Deep Focus Product Specification v1.0.docx',
        'id': 'readiness-focus',
        'title': 'Readiness → Deep Focus',
        'question': 'Can I create and return to one chosen attention window?',
        'prefix': 'DF',
        'first_table': 10,
        'stride': 2,
        'states': ['Select', 'Prepare', 'Enter', 'Sustain', 'Self-direct'],
    },
    {
        'file': 'Transition 004 – Deep Focus → Recovery Product Specification v1.0.docx',
        'id': 'focus-recovery',
        'title': 'Deep Focus → Recovery',
        'question': 'Can I leave effort and choose what comes next?',
        'prefix': 'FR',
        'first_table': 7,
        'stride': 1,
        'states': ['Notice', 'Release', 'Reset', 'Restore', 'Self-direct'],
    },
    {
        'file': 'Transition 005 – Recovery → Presence Product Specification v1.0.docx',
        'source': DELIVERABLES,
        'id': 'recovery-presence',
        'title': 'Recovery → Presence',
        'question': 'Can I arrive in the next real moment?',
        'prefix': 'RP',
        'first_table': 7,
        'stride': 1,
        'states': ['Arrive', 'Orient', 'Participate', 'Return', 'Self-direct'],
    },
    {
        'file': 'Transition 006 – Presence → Wind-down Product Specification v1.0.docx',
        'source': DELIVERABLES,
        'id': 'presence-wind-down',
        'title': 'Presence → Wind-down',
        'question': 'Can I end this part of the day and make room for what follows?',
        'prefix': 'PW',
        'first_table': 7,
        'stride': 1,
        'states': ['Notice', 'Contain', 'Shift', 'Return', 'Self-direct'],
    },
    {
        'file': 'Transition 007 – Wind-down → Release Product Specification v1.0.docx',
        'source': DELIVERABLES,
        'id': 'wind-down-release',
        'title': 'Wind-down → Release',
        'question': 'Can I loosen my hold without requiring resolution?',
        'prefix': 'WR',
        'id_prefix': 'T007-WR',
        'first_table': 7,
        'stride': 1,
        'states': ['Permission', 'Set down', 'Release', 'Return', 'Self-direct'],
    },
]


def normalize(value: str) -> str:
    return ' '.join(value.replace('\u00a0', ' ').split())


def headings(document: Document, prefix: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for paragraph in document.paragraphs:
        text = normalize(paragraph.text)
        match = re.match(rf'^({prefix}-\d{{2}})\s+(.+)$', text)
        if match:
            result[match.group(1)] = match.group(2)
    return result


def authoring_fields(document: Document, prefix: str) -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    current_id: str | None = None
    for paragraph in document.paragraphs:
        text = normalize(paragraph.text)
        heading = re.match(rf'^({prefix}-\d{{2}})\s+(.+)$', text)
        if heading:
            current_id = heading.group(1)
            result[current_id] = {}
            continue
        if not current_id:
            continue
        belief = re.match(r'^Belief under test:\s*(.+)$', text)
        capability = re.match(r'^(?:User outcome|Learning outcome):\s*(.+)$', text)
        if belief:
            result[current_id]['beliefUnderTest'] = belief.group(1)
        if capability:
            result[current_id]['capability'] = capability.group(1)
    return result


def table_fields(document: Document, table_number: int) -> dict[str, str]:
    table = document.tables[table_number - 1]
    fields: dict[str, str] = {}
    for row in table.rows[1:]:
        key = normalize(row.cells[0].text).lower()
        value = normalize(row.cells[1].text)
        fields[key] = value
    return fields


def option_id(index: int) -> str:
    return f'option-{index + 1}'


def build_transition(spec: dict[str, object]) -> dict[str, object]:
    document = Document(Path(spec.get('source', SOURCE)) / str(spec['file']))
    card_titles = headings(document, str(spec['prefix']))
    authored = authoring_fields(document, str(spec['prefix']))
    experiences = []
    for index in range(15):
        source_experience_id = f"{spec['prefix']}-{index + 1:02d}"
        experience_id = f"{spec.get('id_prefix', spec['prefix'])}-{index + 1:02d}"
        fields = table_fields(document, int(spec['first_table']) + index * int(spec['stride']))
        options = [item.strip() for item in fields['response options'].split(' / ')]
        source_material = '\n'.join(fields.values())
        experiences.append({
            'id': experience_id,
            'sourceId': source_experience_id,
            'version': '1.0',
            'semanticVersion': '1.0.0',
            'state': list(spec['states'])[index // 3],
            'title': card_titles[source_experience_id],
            'beliefUnderTest': authored[source_experience_id]['beliefUnderTest'],
            'capability': authored[source_experience_id]['capability'],
            'prompt': fields['prompt'],
            'choicePrompt': fields.get('prediction', fields['prompt']),
            'options': [{'id': option_id(option_index), 'label': option} for option_index, option in enumerate(options)],
            'action': fields['action'],
            'observation': fields['observation'],
            'reflection': fields['reflection'],
            'completion': fields['completion'],
            'fallback': 'Leave it for today. Your place will be kept.',
            'contentHash': hashlib.sha256(source_material.encode()).hexdigest()[:12],
        })
    return {
        'id': spec['id'],
        'title': spec['title'],
        'question': spec['question'],
        'states': spec['states'],
        'experiences': experiences,
    }


def main() -> None:
    transitions = [build_transition(spec) for spec in SPECS]
    target = ROOT / 'src/content/generated.ts'
    target.write_text(
        "import type { Transition } from '@/lib/types';\n"
        "import { validateTransitions } from './validate';\n\n"
        "// Generated from the frozen Transition 001–007 specifications. Do not edit by hand.\n"
        f"export const transitions: Transition[] = {json.dumps(transitions, ensure_ascii=False, indent=2)};\n\n"
        "validateTransitions(transitions);\n",
        encoding='utf-8',
    )
    print(f'Wrote {target} with {sum(len(item["experiences"]) for item in transitions)} experiences.')


if __name__ == '__main__':
    main()
